"""Fase 10: recopilaciones (fragmentos capturados de cualquier materia, ordenables y exportables a PDF y Word)."""
from __future__ import annotations

import io
import json
import sqlite3
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from PIL import Image

from app.errores import ErrorUsuario
from app.services import biblioteca as bib
from app.services import imagenes_fila as imf
from app.services import notas as notas_srv
from app.services import recopilaciones as rec
from app.services import tablas as tab

from test_fase2_excel import configurar, excel_medicina, subir
from test_fase3_pdf import pdf_texto
from test_fase3_pdf import subir as subir_pdf


def material(con: sqlite3.Connection, tmp_path: Path) -> dict:
    """Una ficha con imagen, una nota y un PDF en una materia de prueba."""
    mid = bib.crear_materia(con, "Anatomía", "", "azul", "bone")
    cid = bib.crear_componente(con, mid, None, "Parcial 1", "", ["excel", "notas", "pdf"])
    [tid] = tab.importar(con, configurar(subir(con, tmp_path, cid, excel_medicina(tmp_path / "a.xlsx"))))
    femur = tab.consultar(con, tab.obtener(con, tid), tab.Consulta(q="femur"))[0][0]
    tmp = tmp_path / "rx.png"
    imagen = Image.new("RGB", (400, 200), "white")
    imagen.paste((200, 0, 0), (200, 0, 400, 200))     # mitad derecha roja
    imagen.save(tmp)
    img, _ = imf.agregar(con, femur.id, "rx.png", tmp, "Radiografía AP")
    nota = notas_srv.crear(con, cid, "Resumen", "# Fémur\n\n- **Largo**\n- Fuerte\n\n| A | B |\n|---|---|\n| 1 | 2 |")
    pdf, _ = subir_pdf(con, tmp_path, cid, pdf_texto(tmp_path / "g.pdf"), "Guía")
    return {"fila": femur, "imagen": img, "nota": nota, "pdf": pdf, "tabla": tab.obtener(con, tid)}


@pytest.fixture()
def rcon(datos: Path):
    c = rec.conectar()
    yield c
    c.close()


def test_capturar_ordenar_y_exportar(con, rcon, tmp_path: Path, datos: Path) -> None:
    m = material(con, tmp_path)
    r = rec.activa(rcon, crear_si_falta=True)
    assert r.titulo == rec.TITULO_POR_DEFECTO and rec.activa_id(rcon) == r.id

    rec.agregar_fila(con, rcon, m["fila"].id, r.id)
    clave_notas = next(c["clave"] for c in m["tabla"].columnas if c["nombre"].startswith("Notas"))
    rec.agregar_campo(con, rcon, m["fila"].id, clave_notas, r.id)
    rec.agregar_nota(con, rcon, m["nota"].id, r.id)
    rec.agregar_texto(con, rcon, "  El fémur   es largo  ", "nota", m["nota"].id, "Subrayado", r.id)
    rec.agregar_imagen(con, rcon, "imagen_fila", m["imagen"].id, r.id, recorte="0.5,0,0.5,1")
    rec.agregar_imagen(con, rcon, "pdf", m["pdf"].id, r.id, pagina=2)
    rec.agregar_seccion(rcon, r.id, "Conclusiones")
    rec.agregar_texto_propio(rcon, r.id, "", "Mi **conclusión**")

    r = rec.obtener(rcon, r.id)
    tipos = [f.tipo for f in r.fragmentos]
    assert tipos == ["campos", "texto", "texto", "texto", "imagen", "imagen", "seccion", "texto"]
    ficha, campo, nota, subrayado, recorte, pagina, seccion, propio = r.fragmentos
    assert ficha.titulo == "Fémur" and ["Región", "Miembro inferior"] in ficha.campos
    assert all(valor for _, valor in ficha.campos)                       # solo campos con datos
    assert ficha.origen == "Anatomía › Parcial 1 › Atlas de huesos · Huesos › Fémur" and ficha.origen_url == f"/filas/{m['fila'].id}"
    assert campo.titulo == "Fémur — Notas · Característico" and campo.contenido == "Hueso más largo\ndel cuerpo"
    assert nota.formato == "markdown" and "<table>" in nota.html
    assert subrayado.contenido == "El fémur es largo" and subrayado.titulo == "Subrayado"
    # El recorte es solo la mitad roja, a resolución completa
    con_recorte = Image.open(rec.ruta_imagen(recorte.archivo))
    assert con_recorte.size == (200, 200) and con_recorte.getpixel((100, 100)) == (200, 0, 0)
    assert recorte.titulo == "Radiografía AP (recorte)" and recorte.tamano == "mediana"
    assert pagina.origen.endswith("Guía › página 2") and pagina.origen_url.endswith("?pagina=2") and pagina.tamano == "grande"
    assert propio.origen == "" and propio.formato == "markdown"

    # Copia fija: cambiar el original no cambia la recopilación
    notas_srv.editar(con, m["nota"].id, "Resumen", "otra cosa", "markdown")
    assert rec.obtener_fragmento(rcon, nota.id).contenido.startswith("# Fémur")

    # Mover y ordenar
    rec.mover(rcon, propio.id, -1)
    assert [f.id for f in rec.obtener(rcon, r.id).fragmentos][-2:] == [propio.id, seccion.id]
    rec.mover(rcon, ficha.id, -1)                                         # el primero no sube más
    assert rec.obtener(rcon, r.id).fragmentos[0].id == ficha.id
    ids = [f.id for f in reversed(rec.obtener(rcon, r.id).fragmentos)]
    rec.ordenar(rcon, r.id, ids)
    assert [f.id for f in rec.obtener(rcon, r.id).fragmentos] == ids
    with pytest.raises(ErrorUsuario):
        rec.ordenar(rcon, r.id, ids[1:])

    # Comentarios y tamaño
    rec.editar_fragmento(rcon, recorte.id, "Rojo", "Se ve **claro**", tamano="pequena")
    assert rec.obtener_fragmento(rcon, recorte.id).medidas()[0] == 130   # 200 px: no se agranda hasta los 170 pt
    with pytest.raises(ErrorUsuario):
        rec.editar_fragmento(rcon, subrayado.id, "x", "", contenido="   ")

    # Exportar
    r = rec.obtener(rcon, r.id)
    with pymupdf.open("pdf", rec.exportar_pdf(r)) as doc:
        texto = "".join(p.get_text() for p in doc)
        imagenes = sum(len(p.get_images()) for p in doc)
    for esperado in ("Mi recopilación", "Fémur", "Región:", "El fémur es largo", "Conclusiones", "Se ve claro",
                     "Fuente: Anatomía › Parcial 1", "página 1 de"):
        assert esperado in texto, esperado
    assert imagenes == 2
    word = Document(io.BytesIO(rec.exportar_docx(r)))
    texto_word = "\n".join(p.text for p in word.paragraphs)
    for esperado in ("Mi recopilación", "Región: Miembro inferior", "El fémur es largo", "Conclusiones", "Se ve claro", "Largo"):
        assert esperado in texto_word, esperado
    assert len(word.tables) == 1 and word.tables[0].cell(1, 1).text == "2"
    assert len(word.inline_shapes) == 2

    # Quitar un fragmento borra su imagen copiada; eliminar la recopilación borra todo
    archivo = rec.ruta_imagen(pagina.archivo)
    rec.eliminar_fragmento(rcon, pagina.id)
    assert not archivo.exists()
    restante = rec.ruta_imagen(recorte.archivo)
    rec.eliminar(rcon, r.id)
    assert not restante.exists() and rec.listar(rcon) == [] and rec.activa_id(rcon) is None


def test_validaciones(con, rcon, tmp_path: Path) -> None:
    m = material(con, tmp_path)
    r = rec.crear(rcon, "Caso")
    with pytest.raises(ErrorUsuario, match="nombre"):
        rec.crear(rcon, "   ")
    with pytest.raises(ErrorUsuario, match="subraya"):
        rec.agregar_texto(con, rcon, "  ", "nota", m["nota"].id, "", r.id)
    with pytest.raises(ErrorUsuario, match="pequeño"):
        rec.agregar_imagen(con, rcon, "imagen_fila", m["imagen"].id, r.id, recorte="0.1,0.1,0.001,0.5")
    with pytest.raises(ErrorUsuario, match="páginas"):
        rec.agregar_imagen(con, rcon, "pdf", m["pdf"].id, r.id, pagina=9)
    with pytest.raises(ErrorUsuario, match="vacío"):
        rec.agregar_campo(con, rcon, m["fila"].id, "no-existe", r.id)
    with pytest.raises(ErrorUsuario):
        rec.ruta_imagen("../biblioteca.db")
    assert rec.leer_recorte("0.9,0.9,0.5,0.5") == pytest.approx((0.9, 0.9, 0.1, 0.1))   # se ajusta a la imagen
    assert rec.nombre_archivo(rec.Recopilacion(1, 'Caso: "A/B"?', "", "", ""), "pdf") == "Caso A B.pdf"


def test_recopilacion_desde_la_web(cliente, tmp_path: Path, datos: Path) -> None:
    from app import db
    con = db.conectar()
    m = material(con, tmp_path)
    con.close()
    hx = {"HX-Request": "true"}

    # Sin recopilaciones: agregar crea «Mi recopilación» y avisa
    r = cliente.post("/recopilacion/agregar", data={"fuente": "fila", "ref_id": m["fila"].id}, headers=hx)
    assert r.status_code == 204
    disparo = json.loads(r.headers["HX-Trigger"])
    assert "Mi recopilación" in disparo["aviso"]["mensaje"] and disparo["recopilacion-cambio"]
    r = cliente.post("/recopilacion/agregar", data={"fuente": "texto:fila", "ref_id": m["fila"].id, "texto": "Miembro inferior",
                                                    "titulo": "Fémur — Región"}, headers=hx)
    assert r.status_code == 204
    r = cliente.post("/recopilacion/agregar", data={"fuente": "pdf", "ref_id": m["pdf"].id, "pagina": 1,
                                                    "recorte": "0,0,1,0.3"}, headers=hx)
    assert r.status_code == 204 and "recorte" in json.loads(r.headers["HX-Trigger"])["aviso"]["mensaje"]

    assert "Mi recopilación" in cliente.get("/recopilacion/contador").text
    lista = cliente.get("/recopilaciones")
    assert lista.status_code == 200 and "Activa" in lista.text
    rcon = rec.conectar()
    [recop] = rec.listar(rcon)
    pagina = cliente.get(f"/recopilaciones/{recop.id}")
    assert pagina.status_code == 200 and "Fémur" in pagina.text and "Miembro inferior" in pagina.text

    # Ventana de recorte (imagen y página de PDF) y la imagen de la página
    assert "recorte(" in cliente.get(f"/recopilacion/recortar?fuente=imagen_fila&ref_id={m['imagen'].id}").text
    assert "de 3" in cliente.get(f"/recopilacion/recortar?fuente=pdf&ref_id={m['pdf'].id}&pagina=2").text
    png = cliente.get(f"/recopilacion/pdf/{m['pdf'].id}/pagina?n=2")
    assert png.headers["content-type"] == "image/png" and Image.open(io.BytesIO(png.content)).width > 500

    # Mover, editar, agregar sección y quitar (la lista se devuelve para reemplazarla)
    frag = rec.obtener(rcon, recop.id).fragmentos
    r = cliente.post(f"/fragmentos/{frag[2].id}/mover", data={"paso": -1}, headers=hx)
    assert r.status_code == 200 and [f.id for f in rec.obtener(rcon, recop.id).fragmentos][1] == frag[2].id
    r = cliente.post(f"/fragmentos/{frag[0].id}", data={"titulo": "Fémur (ficha)", "comentario": "Clave para el parcial"}, headers=hx)
    assert "Clave para el parcial" in r.text and json.loads(r.headers["HX-Trigger"])["cerrar-modal"]
    assert "Nueva sección" in cliente.get(f"/recopilaciones/{recop.id}/texto?tipo=seccion").text
    cliente.post(f"/recopilaciones/{recop.id}/seccion", data={"titulo": "Final"}, headers=hx)
    cliente.post(f"/recopilaciones/{recop.id}/orden", json={"ids": [f.id for f in rec.obtener(rcon, recop.id).fragmentos][::-1]})
    assert rec.obtener(rcon, recop.id).fragmentos[0].titulo == "Final"

    # Vista previa y descargas
    vista = cliente.get(f"/recopilaciones/{recop.id}/vista-previa")
    assert vista.status_code == 200 and "documento-recopilacion" in vista.text and "Clave para el parcial" in vista.text
    pdf = cliente.get(f"/recopilaciones/{recop.id}/pdf")
    assert pdf.headers["content-type"] == "application/pdf" and "Mi%20recopilaci%C3%B3n.pdf" in pdf.headers["content-disposition"]
    word = cliente.get(f"/recopilaciones/{recop.id}/word")
    assert word.status_code == 200 and word.content[:2] == b"PK"
    nombre = rec.obtener(rcon, recop.id).fragmentos[1].archivo or rec.obtener(rcon, recop.id).fragmentos[2].archivo
    assert cliente.get(f"/recopilaciones/imagenes/{nombre}").status_code == 200

    # Varias recopilaciones: crear una nueva la vuelve activa; se puede volver a la anterior y eliminar
    r = cliente.post("/recopilaciones", data={"titulo": "Caso 2"}, headers=hx)
    assert r.headers["HX-Redirect"].startswith("/recopilaciones/")
    assert rec.activa(rcon).titulo == "Caso 2"
    cliente.post(f"/recopilaciones/{recop.id}/activar", headers=hx)
    assert rec.activa_id(rcon) == recop.id
    cliente.post(f"/recopilaciones/{recop.id}/eliminar", headers=hx)
    assert [x.titulo for x in rec.listar(rcon)] == ["Caso 2"]
    assert cliente.get(f"/recopilaciones/{recop.id}").status_code == 404
    rcon.close()


def test_pagina_de_ficha_tiene_botones_de_captura(cliente, tmp_path: Path) -> None:
    from app import db
    con = db.conectar()
    m = material(con, tmp_path)
    con.close()
    html = cliente.get(f"/filas/{m['fila'].id}").text
    assert 'data-capturar="fila:' in html and '"fuente": "campo"' in html and "Recopilar ficha" in html
    assert 'data-capturar="nota:' in cliente.get(f"/notas/{m['nota'].id}").text
    assert "fuente=pdf" in cliente.get(f"/pdfs/{m['pdf'].id}").text


def test_agregar_varios_seleccionados(cliente, tmp_path: Path) -> None:
    from app import db
    con = db.conectar()
    m = material(con, tmp_path)
    con.close()
    elementos = [{"fuente": "nota", "ref_id": m["nota"].id, "titulo": "Resumen"},
                 {"fuente": "fila", "ref_id": m["fila"].id, "titulo": "Fémur"},
                 {"fuente": "imagen_fila", "ref_id": m["imagen"].id, "titulo": "Radiografía"},
                 {"fuente": "pdf", "ref_id": m["pdf"].id, "pagina": 3, "titulo": "Guía"},
                 {"fuente": "fila", "ref_id": 99999, "titulo": "Borrada"}]
    r = cliente.post("/recopilacion/agregar-varios", json={"elementos": elementos}, headers={"HX-Request": "true"})
    assert r.status_code == 200
    aviso = json.loads(r.headers["HX-Trigger"])["aviso"]
    assert aviso["tipo"] == "aviso" and "agregaron 4 elementos" in aviso["mensaje"] and "Borrada" in aviso["mensaje"]
    rcon = rec.conectar()
    [recop] = rec.listar(rcon)
    assert [f.tipo for f in rec.obtener(rcon, recop.id).fragmentos] == ["texto", "campos", "imagen", "imagen"]
    rcon.close()

    # Las listas muestran las casillas
    assert '"fuente": "fila"' in cliente.get(f"/tablas/{m['tabla'].id}").text.replace("&#34;", '"')
    assert "lista-notas" in cliente.get(f"/componentes/{m['tabla'].componente_id}").text


def test_explorar_permite_recopilar(cliente, tmp_path: Path) -> None:
    from app import db
    con = db.conectar()
    m = material(con, tmp_path)
    comp = bib.obtener_componente(con, m["tabla"].componente_id)
    con.close()
    base = f"/materias/{comp.materia_id}/explorar?c={comp.id}"
    nota = cliente.get(f"{base}&ver=nota-{m['nota'].id}").text
    assert 'data-capturar="nota:' in nota and '"fuente": "nota"' in nota
    tabla = cliente.get(f"{base}&ver=tabla-{m['tabla'].id}").text.replace("&#34;", '"')
    assert '"fuente": "fila"' in tabla and "Seleccionar todas las filas mostradas" in tabla
    assert "fuente=pdf" in cliente.get(f"{base}&ver=pdf-{m['pdf'].id}").text
