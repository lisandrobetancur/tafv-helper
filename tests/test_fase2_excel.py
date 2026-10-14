"""Fase 2: lectura de Excel/CSV, importación, consulta, imágenes por fila y reemplazo."""
from __future__ import annotations

import io
import sqlite3
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from PIL import Image

from app.errores import ErrorUsuario
from app.services import biblioteca as bib
from app.services import hojas as lector
from app.services import imagenes_fila as imf
from app.services import importacion as imp
from app.services import papelera as pap
from app.services import tablas as tab


# ------------------------------------------------------------------ utilidades
def excel_medicina(ruta: Path, filas_extra: list[list] | None = None) -> Path:
    """Excel con lo difícil: título combinado, encabezado de 2 filas, fechas, vacíos, texto largo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Huesos"
    ws["A1"] = "ATLAS DE ANATOMÍA"
    ws.merge_cells("A1:D1")
    ws.append(["Nombre", "Región", "Fecha repaso", "Notas"])            # fila 2: encabezado
    ws.append(["Nombre", "Región", "Fecha repaso", "Característico"])   # fila 3: subencabezado
    filas = [
        ["Fémur", "Miembro inferior", datetime(2025, 7, 6), "Hueso más largo\ndel cuerpo"],
        ["Húmero", "Miembro superior", None, ""],
        [None, None, None, None],                                         # fila vacía
        ["Tibia", "Miembro inferior", datetime(2025, 3, 1), "x" * 500],
    ] + (filas_extra or [])
    for f in filas:
        ws.append(f)
    ws["C4"].number_format = "d-m"        # "6-7" que Excel convirtió en fecha
    ws["C7"].number_format = "dd/mm/yyyy"
    ws.append(["Columna combinada", "Pelvis", None, None])
    ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=3)
    wb.create_sheet("Vacía")
    otra = wb.create_sheet("Músculos")
    otra.append(["Músculo", "Inervación"])
    otra.append(["Bíceps", "Musculocutáneo"])
    wb.save(ruta)
    return ruta


def png(color: str = "red", tamano: tuple[int, int] = (40, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", tamano, color).save(buf, "PNG")
    return buf.getvalue()


def preparar_componente(con: sqlite3.Connection) -> int:
    mid = bib.crear_materia(con, "Anatomía", "", "azul", "bone")
    return bib.crear_componente(con, mid, None, "Parcial 1", "", ["excel"])


def subir(con, tmp_path: Path, cid: int, ruta: Path, nombre: str = "Atlas de huesos", tabla_id=None) -> imp.Estado:
    temporal = tmp_path / f"subida-{ruta.name}"
    temporal.write_bytes(ruta.read_bytes())
    return imp.crear(con, cid, ruta.name, temporal, nombre, tabla_id=tabla_id)


def configurar(estado: imp.Estado, indice: int | None = None) -> imp.Estado:
    hojas = imp.hojas(estado)
    for i in ([indice] if indice is not None else estado.seleccion):
        estado.poner_config(i, imp.config_sugerida(estado, i, hojas[i]))
    imp.guardar(estado)
    return estado


# ------------------------------------------------------------------- lectura
def test_lectura_excel_dificil(tmp_path: Path) -> None:
    hojas = lector.leer_archivo(excel_medicina(tmp_path / "a.xlsx"))
    assert [h.nombre for h in hojas] == ["Huesos", "Vacía", "Músculos"]
    h = hojas[0]
    assert h.filas[0][:4] == ["ATLAS DE ANATOMÍA"] * 4             # combinada: se repite
    assert h.filas[3][2] == "6-7"                                     # fecha "como se ve"
    assert h.filas[6][2] == "01/03/2025"
    assert h.filas[3][3] == "Hueso más largo\ndel cuerpo"             # saltos de línea intactos
    assert h.filas[7][1:3] == ["Pelvis", "Pelvis"]                    # combinada horizontal
    assert lector.detectar_encabezado(h) == 2
    assert lector.parece_subencabezado(h, 2)
    cols = lector.columnas(h, 2, dos_filas=True)
    assert [c.nombre for c in cols] == ["Nombre", "Región", "Fecha repaso", "Notas · Característico"]
    regs = lector.registros(h, 2, True, [c.indice for c in cols])
    assert [r[1][0] for r in regs] == ["Fémur", "Húmero", "Tibia", "Columna combinada"]  # sin la vacía
    assert regs[0][0] == 4                                            # número de fila original


def test_lectura_csv_latin1_y_punto_y_coma(tmp_path: Path) -> None:
    ruta = tmp_path / "farmacos.csv"
    ruta.write_bytes("Fármaco;Dosis\nIbuprofeno;400 mg\nParacetamol;1 g\n".encode("cp1252"))
    [h] = lector.leer_archivo(ruta)
    assert h.nombre == "farmacos" and h.filas[0] == ["Fármaco", "Dosis"] and h.filas[2] == ["Paracetamol", "1 g"]


def test_formatos_no_compatibles(tmp_path: Path) -> None:
    (tmp_path / "viejo.xls").write_bytes(b"x")
    with pytest.raises(ErrorUsuario, match=".xlsx"):
        lector.leer_archivo(tmp_path / "viejo.xls")
    (tmp_path / "roto.xlsx").write_bytes(b"no es un excel")
    with pytest.raises(ErrorUsuario, match="No se pudo abrir"):
        lector.leer_archivo(tmp_path / "roto.xlsx")


@pytest.mark.parametrize("valor, formato, esperado", [
    (3.0, "General", "3"), (0.125, "0%", "12.5%"), (1234.5, "General", "1234.5"), (True, "General", "Sí"),
    (datetime(2025, 7, 6, 14, 30), "dd/mm/yyyy hh:mm", "06/07/2025 14:30"), (None, "General", ""),
])
def test_texto_celda(valor, formato, esperado) -> None:
    assert lector.texto_celda(valor, formato) == esperado


# ---------------------------------------------------------------- importación
def test_importar_varias_hojas(con, tmp_path: Path, datos: Path) -> None:
    cid = preparar_componente(con)
    estado = subir(con, tmp_path, cid, excel_medicina(tmp_path / "atlas_huesos.xlsx"))
    assert estado.nombre == "Atlas de huesos"
    estado.seleccion = [0, 2]
    configurar(estado)
    assert estado.config_hoja(0).dos_filas and estado.config_hoja(0).fila_encabezado == 2
    ids = tab.importar(con, estado)

    tablas = tab.listar(con, cid)
    assert [t.nombre for t in tablas] == ["Atlas de huesos · Huesos", "Atlas de huesos · Músculos"]
    assert tablas[0].num_filas == 4 and tablas[1].num_filas == 1
    # el archivo se guarda con el nombre descriptivo, en la carpeta del componente, una sola vez
    assert (datos / "archivos" / "anatomia" / "parcial-1" / "excel" / "Atlas de huesos.xlsx").exists()
    assert tablas[0].archivo == tablas[1].archivo == "excel/Atlas de huesos.xlsx"
    assert not (datos / "cache" / "importaciones" / estado.token).exists()   # limpieza

    fila = tab.consultar(con, tab.obtener(con, ids[0]), tab.Consulta())[0][0]
    assert fila.titulo == "Fémur" and fila.fila_origen == 4
    assert bib.contar_elementos_componente(con, cid) == 2


def test_duplicado_se_avisa(con, tmp_path: Path) -> None:
    cid = preparar_componente(con)
    ruta = excel_medicina(tmp_path / "a.xlsx")
    tab.importar(con, configurar(subir(con, tmp_path, cid, ruta)))
    assert subir(con, tmp_path, cid, ruta).duplicado_de == "Atlas de huesos · Huesos"


def test_nombre_descriptivo_obligatorio(con, tmp_path: Path) -> None:
    cid = preparar_componente(con)
    with pytest.raises(ErrorUsuario, match="nombre descriptivo"):
        subir(con, tmp_path, cid, excel_medicina(tmp_path / "a.xlsx"), nombre="  ")


# ------------------------------------------------------------------- consulta
def test_consulta_sin_tildes_filtros_y_orden(con, tmp_path: Path) -> None:
    cid = preparar_componente(con)
    [tid] = tab.importar(con, configurar(subir(con, tmp_path, cid, excel_medicina(tmp_path / "a.xlsx"))))
    t = tab.obtener(con, tid)
    clave_region = t.claves[1]
    titulos = lambda c: [f.titulo for f in tab.consultar(con, t, c)[0]]
    assert titulos(tab.Consulta(q="FEMUR")) == ["Fémur"]
    assert titulos(tab.Consulta(q="miembro inferior")) == ["Fémur", "Tibia"]
    assert titulos(tab.Consulta(filtros={clave_region: "superior"})) == ["Húmero"]
    assert titulos(tab.Consulta(orden=t.claves[0], desc=True))[0] == "Tibia"
    assert titulos(tab.Consulta(q="100%")) == []                       # comodines escapados
    filas, total = tab.consultar(con, t, tab.Consulta(por_pagina=25, pagina=9))
    assert total == 4 and len(filas) == 4                              # página fuera de rango -> última


def test_consulta_rapida_con_miles_de_filas(con, tmp_path: Path) -> None:
    import time
    extra = [[f"Estructura {i}", f"Región {i % 50}", None, f"Descripción larga número {i}"] for i in range(5000)]
    cid = preparar_componente(con)
    [tid] = tab.importar(con, configurar(subir(con, tmp_path, cid, excel_medicina(tmp_path / "g.xlsx", extra))))
    t = tab.obtener(con, tid)
    inicio = time.perf_counter()
    filas, total = tab.consultar(con, t, tab.Consulta(filtros={t.claves[1]: "región 7"}, orden=t.claves[0], pagina=2))
    # orden natural: "Estructura 9" antes que "Estructura 10"
    primeros = [f.titulo for f in tab.consultar(con, t, tab.Consulta(q="estructura", orden=t.claves[0], por_pagina=25))[0]][:12]
    assert primeros[:3] == ["Estructura 0", "Estructura 1", "Estructura 2"] and "Estructura 10" == primeros[10]
    assert time.perf_counter() - inicio < 0.3
    assert total == 100 and len(filas) == 50


# ----------------------------------------------------------- imágenes por fila
def test_imagenes_de_fila(con, tmp_path: Path, datos: Path) -> None:
    cid = preparar_componente(con)
    [tid] = tab.importar(con, configurar(subir(con, tmp_path, cid, excel_medicina(tmp_path / "a.xlsx"))))
    femur = tab.consultar(con, tab.obtener(con, tid), tab.Consulta(q="femur"))[0][0]
    tmp = tmp_path / "rx.png"
    tmp.write_bytes(png())
    img, aviso = imf.agregar(con, femur.id, "IMG_2031.PNG", tmp, "Radiografía AP", "Vista anterior")
    assert aviso is None and img.titulo == "Radiografía AP" and (img.ancho, img.alto) == (40, 30)
    assert img.archivo == "imagenes-filas/Fémur - Radiografía AP.png"     # nombre legible en disco
    assert (datos / "archivos" / "anatomia" / "parcial-1" / img.archivo).exists()

    assert imf.agregar(con, femur.id, "otra.png", tmp, "Repetida")[0] is None  # misma imagen, misma ficha
    with pytest.raises(ErrorUsuario, match="nombre descriptivo"):
        imf.agregar(con, femur.id, "x.png", tmp, " ")
    malo = tmp_path / "malo.png"
    malo.write_bytes(b"no soy imagen")
    with pytest.raises(ErrorUsuario, match="no es una imagen"):
        imf.agregar(con, femur.id, "malo.png", malo, "Mala")

    tmp2 = tmp_path / "b.png"
    tmp2.write_bytes(png("blue"))
    img2, _ = imf.agregar(con, femur.id, "b.png", tmp2, "Corte sagital")
    imf.reordenar(con, femur.id, [img2.id, img.id])
    assert [i.titulo for i in imf.listar(con, femur.id)] == ["Corte sagital", "Radiografía AP"]

    editada = imf.editar(con, img.id, "Rx anteroposterior", "Nuevo pie")
    assert editada.archivo == "imagenes-filas/Fémur - Rx anteroposterior.png" and editada.pie == "Nuevo pie"
    assert (datos / "archivos" / "anatomia" / "parcial-1" / editada.archivo).exists()

    imf.enviar_a_papelera(con, img2.id)
    assert [i.titulo for i in imf.listar(con, femur.id)] == ["Rx anteroposterior"]
    [e] = pap.listar(con)
    pap.restaurar(con, e.id)
    assert len(imf.listar(con, femur.id)) == 2


# ------------------------------------------------------------------ reemplazo
def test_reemplazar_excel_conserva_imagenes(con, tmp_path: Path, datos: Path) -> None:
    cid = preparar_componente(con)
    [tid] = tab.importar(con, configurar(subir(con, tmp_path, cid, excel_medicina(tmp_path / "v1.xlsx"))))
    t = tab.obtener(con, tid)
    filas = {f.titulo: f for f in tab.consultar(con, t, tab.Consulta())[0]}
    tmp = tmp_path / "i.png"
    tmp.write_bytes(png())
    imf.agregar(con, filas["Fémur"].id, "i.png", tmp, "Rx fémur")
    tmp.write_bytes(png("green"))
    imf.agregar(con, filas["Húmero"].id, "i.png", tmp, "Rx húmero")

    # Versión 2: Fémur cambia, Húmero desaparece (tiene imagen), Tibia desaparece (sin imagen), Radio es nueva
    wb = Workbook()
    ws = wb.active
    ws.title = "Huesos"
    ws.append(["Nombre", "Región", "Notas"])
    ws.append(["FEMUR", "Muslo", "Actualizado"])
    ws.append(["Radio", "Antebrazo", ""])
    ws.append(["Columna combinada", "Pelvis", ""])
    wb.save(tmp_path / "v2.xlsx")
    estado = subir(con, tmp_path, cid, tmp_path / "v2.xlsx", nombre=t.nombre, tabla_id=tid)
    estado.poner_config(0, imp.config_sugerida(estado, 0, imp.hojas(estado)[0], vars(t) | {"columnas": t.columnas}))
    imp.guardar(estado)
    previa = tab.previsualizar_reemplazo(con, estado)
    resumen = tab.reemplazar(con, estado)
    assert previa == resumen
    assert (resumen.actualizadas, resumen.nuevas, resumen.conservadas, resumen.eliminadas) == (2, 1, 1, 1)

    t = tab.obtener(con, tid)
    todas = {f.titulo: f for f in tab.consultar(con, t, tab.Consulta())[0]}
    assert set(todas) == {"FEMUR", "Radio", "Columna combinada", "Húmero"}
    assert todas["FEMUR"].id == filas["Fémur"].id and todas["FEMUR"].num_imagenes == 1  # misma ficha, misma imagen
    assert todas["Húmero"].huerfana and todas["Húmero"].num_imagenes == 1
    assert todas["FEMUR"].valores[t.claves[1]] == "Muslo"
    # misma clave para la columna "Región" en ambas versiones; "Notas" también se reconoce por nombre
    assert t.nombre_columna(t.claves[1]) == "Región"
    # el Excel anterior queda recuperable en la papelera
    assert any(e.tipo == "archivo" for e in pap.listar(con))

    assert tab.eliminar_huerfanas(con, tid) == 1
    assert "Húmero" not in {f.titulo for f in tab.consultar(con, tab.obtener(con, tid), tab.Consulta())[0]}


def test_actualizar_respeta_el_nombre_aunque_haya_varias_hojas(con, tmp_path: Path) -> None:
    cid = preparar_componente(con)
    ruta = excel_medicina(tmp_path / "a.xlsx")
    [tid] = tab.importar(con, configurar(subir(con, tmp_path, cid, ruta)))
    t = tab.obtener(con, tid)
    estado = subir(con, tmp_path, cid, ruta, nombre=t.nombre, tabla_id=tid)
    cfg = imp.config_sugerida(estado, 0, imp.hojas(estado)[0], {"columnas": t.columnas, "columna_titulo": t.columna_titulo,
                                                                 "columna_clave": t.columna_clave})
    assert cfg.nombre_tabla == t.nombre == "Atlas de huesos · Huesos"


# ------------------------------------------------------------------- papelera
def test_tabla_a_papelera_y_vaciar_borra_archivos(con, tmp_path: Path, datos: Path) -> None:
    cid = preparar_componente(con)
    [tid] = tab.importar(con, configurar(subir(con, tmp_path, cid, excel_medicina(tmp_path / "a.xlsx"))))
    fila = tab.consultar(con, tab.obtener(con, tid), tab.Consulta())[0][0]
    tmp = tmp_path / "i.png"
    tmp.write_bytes(png())
    img, _ = imf.agregar(con, fila.id, "i.png", tmp, "Rx")
    carpeta = datos / "archivos" / "anatomia" / "parcial-1"

    pap.enviar_tabla(con, tid)
    assert tab.listar(con, cid) == [] and bib.contar_elementos_componente(con, cid) == 0
    pap.restaurar(con, pap.listar(con)[0].id)
    assert len(tab.listar(con, cid)) == 1

    pap.enviar_tabla(con, tid)
    pap.vaciar(con)
    assert not (carpeta / img.archivo).exists() and not (carpeta / "excel" / "Atlas de huesos.xlsx").exists()
    assert con.execute("SELECT COUNT(*) FROM filas").fetchone()[0] == 0


def test_titulo_sugerido_con_columna_agrupadora(tmp_path: Path) -> None:
    """«Sistema» (combinada, se repite) no es el título aunque contenga «tema»; «Nombre» sí."""
    wb = Workbook()
    ws = wb.active
    ws.append([])
    ws.append(["Sistema", "Grupo", "Nombre", "Nombre Cientifico", "Accion"])
    ws.append(["Respiratorio", "Asma", "Totumo", "Crescentia cujete", "Broncodilatador"])
    ws.append([None, "Asma", "Yarumo", "Cecropia peltata", "Broncodilatador\nAntiinflamatorio"])
    ws.append([None, "Epoc", "Borraja", "Borago officinalis", "Mucolítico"])
    ws.merge_cells("A3:A5")
    wb.save(tmp_path / "fito.xlsx")
    hoja = lector.leer_archivo(tmp_path / "fito.xlsx")[0]
    fila = lector.detectar_encabezado(hoja)
    cols = lector.columnas(hoja, fila)
    assert lector.columna_titulo_sugerida(cols) == 2
    filas = lector.registros(hoja, fila, False, [c.indice for c in cols])
    assert [f[1][0] for f in filas] == ["Respiratorio"] * 3          # la combinada se repite en cada fila
    assert filas[1][1][4] == "Broncodilatador\nAntiinflamatorio"     # saltos de línea intactos
    sin_preferidas = [lector.Columna(0, "A", "Sistema", ["Resp", "Resp"], 2), lector.Columna(1, "B", "Columna B", ["a", "b"], 2)]
    assert lector.columna_titulo_sugerida(sin_preferidas) == 1       # la primera sin valores repetidos
