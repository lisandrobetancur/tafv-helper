import hashlib
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import main
from app.excel_repo import ExcelRepo, cell_text

PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 32


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(main, "DATA_DIR", tmp_path)  # adjuntos aislados del proyecto
    with TestClient(main.app) as c:
        yield c


def key_of(client, label):
    fields = client.get("/api/meta").json()["fields"]
    return next(f["key"] for f in fields if f["label"] == label)


def test_records_and_filters(client):
    all_ = client.get("/api/records").json()["records"]
    assert len(all_) > 30
    med = key_of(client, "Medicamento")
    lyc = client.get("/api/records", params={"medicamento": "lycop"}).json()["records"]
    assert [r[med].strip() for r in lyc] == ["Lycopodium"]
    assert client.get("/api/records", params={"sintomas": "MIEDO"}).json()["records"]  # sin acentos/mayúsculas
    both = client.get("/api/records", params={"medicamento": "lyc", "sintomas": "MIEDO al fracaso"}).json()["records"]
    assert both and all("lyc" in r[med].lower() for r in both)
    assert client.get("/api/suggest", params={"field": "sintomas", "q": "ans"}).json()
    assert client.get("/api/suggest", params={"field": "nope", "q": "x"}).status_code == 400


def test_filtro_otros_sintomas_y_columna_oculta_en_filtros(client):
    meta = client.get("/api/meta").json()
    assert [f["key"] for f in meta["filters"]] == ["medicamento", "sintomas", "otros_sintomas"]
    hidden = [f["label"] for f in meta["fields"] if f["hide_in_filters"]]
    assert hidden == ["Otros Sintomas · Caracteristico"]  # solo se oculta en la tabla de filtros

    med = key_of(client, "Medicamento")
    frase = "lateralidad derecha y progresion"
    en_otros = client.get("/api/records", params={"otros_sintomas": frase}).json()["records"]
    en_sintomas = client.get("/api/records", params={"sintomas": frase}).json()["records"]
    assert [r[med].strip() for r in en_otros] == ["Lycopodium"]
    assert en_sintomas == []  # "Síntomas" ya no busca en "Otros síntomas"
    assert client.get("/api/suggest", params={"field": "otros_sintomas", "q": "lateralidad"}).json()
    # la ficha (API de registros) sigue trayendo el valor completo
    assert all("otros_sintomas_caracteristico" in r for r in client.get("/api/records").json()["records"])


def test_regla_se_muestra_tal_cual(client):
    med, regla = key_of(client, "Medicamento"), key_of(client, "Regla de preparación")
    reglas = {r[med].strip(): r[regla] for r in client.get("/api/records").json()["records"]}
    assert reglas["Sulfur"] == "6-7"


def test_solo_columnas_b_a_o(client):
    meta = client.get("/api/meta").json()
    assert meta["columns"] == "B:O"
    letters = [f["letter"] for f in meta["fields"]]
    assert letters[0] == "B" and letters[-1] == "O" and len(letters) == 14
    assert all(60 <= f["width"] <= 360 for f in meta["fields"])


def test_formulacion_page(client):
    html = client.get("/formulacion").text
    assert "Dilución" in html and "Decimal" in html


def test_attachments_and_excel_untouched(client):
    excel = Path(main.EXCEL_PATH)
    before = hashlib.md5(excel.read_bytes()).hexdigest()
    rid = client.get("/api/records").json()["records"][0]["id"]

    bad = client.post(f"/api/records/{rid}/attachments", files={"files": ("x.png", b"no soy imagen")})
    assert bad.json()["errors"] and not bad.json()["added"]

    ok = client.post(f"/api/records/{rid}/attachments", files={"files": ("../../x.png", PNG)}).json()
    assert ok["added"] and ok["added"][0]["name"] == "x.png"
    url = ok["added"][0]["url"]
    assert client.get(url).content == PNG
    assert client.delete(url.replace("/file", "")).status_code == 200
    assert client.get(url).status_code == 404
    assert client.post("/api/records/nope/attachments", files={"files": ("a.png", PNG)}).status_code == 404

    assert hashlib.md5(excel.read_bytes()).hexdigest() == before


def test_valores_identicos_al_excel(client):
    """Cada celda de texto del Excel llega a la API sin ningún cambio (espacios y saltos incluidos)."""
    repo = client.app.state.repo
    ws = openpyxl.load_workbook(main.EXCEL_PATH, data_only=True)["Estudio homeopatia"]
    records = {r["row"]: r for r in client.get("/api/records").json()["records"]}
    checked = 0
    for row, rec in records.items():
        for f in repo.fields:
            raw = ws.cell(row, f["col"]).value
            if isinstance(raw, str):
                assert rec[f["key"]] == raw, (row, f["key"])
                checked += 1
    assert checked > 250


def test_fechas_de_excel_se_muestran_como_texto_digitado():
    from datetime import datetime
    from types import SimpleNamespace

    def cell(v, f): return SimpleNamespace(value=v, number_format=f)
    d = datetime(2025, 7, 6)
    assert cell_text(cell(d, r"d\-m")) == "6-7"
    assert cell_text(cell(d, "d/m")) == "6/7"
    assert cell_text(cell(d, "m/d/yy")) == "7/6/25"
    assert cell_text(cell(d, "General")) == "6-7"
    assert cell_text(cell(7.0, "General")) == "7"


def _make_xlsx(path, header1, rows, header2=None, start_col=2):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Estudio homeopatia"
    ws.append([None] * (start_col - 1) + header1)
    if header2:
        ws.append([None] * (start_col - 1) + header2)
    for r in rows:
        ws.append([None] * (start_col - 1) + r)
    wb.create_sheet("Formulación").append(["x"])
    wb.save(path)


def test_campos_dinamicos_menos_columnas(tmp_path):
    f = tmp_path / "a.xlsx"
    _make_xlsx(f, ["Medicamento", "Dosis"], [["Arnica", "30CH"], ["Nux", "200CH"]])
    repo = ExcelRepo(f, "B:O")  # el rango pide hasta O pero el archivo solo tiene B:C
    assert [x["label"] for x in repo.fields] == ["Medicamento", "Dosis"]
    assert len(repo.records) == 2
    assert [x["key"] for x in repo.filters] == ["medicamento"]  # sin síntomas: solo ese filtro
    assert not any(x["hide_in_filters"] for x in repo.fields)
    assert repo.search({"medicamento": "arn"})[0]["dosis"] == "30CH"


def test_campos_dinamicos_mas_columnas_y_sin_encabezados_conocidos(tmp_path):
    f = tmp_path / "b.xlsx"
    headers = [f"Campo {i}" for i in range(1, 21)]
    _make_xlsx(f, headers, [[f"v{i}" for i in range(1, 21)]])
    assert len(ExcelRepo(f, "B:O").fields) == 14           # respeta el rango pedido
    repo = ExcelRepo(f, "B:")                              # rango abierto: todas
    assert len(repo.fields) == 20
    assert [x["key"] for x in repo.filters] == ["buscar"]  # sin encabezados conocidos: búsqueda general
    assert repo.search({"buscar": "v20"})


def test_subencabezado_y_encabezado_compuesto(tmp_path):
    f = tmp_path / "c.xlsx"
    _make_xlsx(f, ["Medicamento", "Otros Sintomas"], [["Arnica", "dolor"]],
               header2=["Medicamento", "Caracteristico"])
    repo = ExcelRepo(f, "B:O")
    assert [x["label"] for x in repo.fields] == ["Medicamento", "Otros Sintomas · Caracteristico"]
    assert repo.records[0]["row"] == 3
    assert [x["key"] for x in repo.filters] == ["medicamento", "otros_sintomas"]
    assert repo.fields[1]["hide_in_filters"]
    assert repo.search({"otros_sintomas": "dolor"})
