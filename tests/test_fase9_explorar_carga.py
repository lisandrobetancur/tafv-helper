"""Fase 9: pestaña «Explorar» de la materia (con carruseles) y herramientas de carga (tools/carga/)."""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XlImagen
from openpyxl.styles import Font, PatternFill
from PIL import Image

from app.services import biblioteca as bib
from app.services import imagenes as img_srv
from app.services import notas as notas_srv

RAIZ = Path(__file__).resolve().parents[1]
CARGA = RAIZ / "tools" / "carga"


# ------------------------------------------------------------------ Explorar
def test_explorar_indice_lectura_y_navegacion(cliente, con, tmp_path: Path) -> None:
    mid = bib.crear_materia(con, "Oxigenoterapia", "", "azul", "wind")
    hbo = bib.crear_componente(con, mid, None, "HBO", "", ["imagenes", "notas"])
    grupo = bib.crear_componente(con, mid, None, "Ozono", "", [])
    hijo = bib.crear_componente(con, mid, grupo, "Vías", "", ["notas"])
    notas_srv.crear(con, hbo, "Protocolo - HBO", "| Sección | Presión |\n|---|---|\n| 1 | 5 p.s.i |")
    notas_srv.crear(con, hijo, "Vías sistémicas", "- Ozono endovenoso")
    for n in (1, 2):
        ruta = tmp_path / f"i{n}.png"
        Image.new("RGB", (40, 30), "white").save(ruta)
        img_srv.agregar(con, hbo, ruta.name, ruta, f"Lámina {n}", f"Texto de la lámina {n}")

    r = cliente.get(f"/materias/{mid}/explorar")
    assert r.status_code == 200
    html = r.text
    assert "Protocolo - HBO" in html and "5 p.s.i" in html            # nota desplegada (tabla Markdown)
    assert 'x-data="carrusel(' in html and "mina 2" in html          # carrusel con las dos imágenes
    assert f"explorar?c={hijo}" in html                               # índice con el subcomponente
    assert "Componentes y carga" in html                              # pestañas de la materia

    r = cliente.get(f"/materias/{mid}/explorar?c={hijo}")
    assert "Ozono" in r.text and "Ozono endovenoso" in r.text
    assert f"explorar?c={hbo}" in r.text                              # «Anterior» salta el grupo vacío

    vacia = bib.crear_materia(con, "Psicoterapia", "", "rojo", "brain")
    assert "Aún no hay nada para explorar" in cliente.get(f"/materias/{vacia}/explorar").text
    assert f"/materias/{mid}/explorar?c={hbo}" in cliente.get(f"/componentes/{hbo}").text   # botón Explorar


# ------------------------------------------------------------------ Herramientas de carga
def _libro(ruta: Path, png: Path) -> None:
    wb = openpyxl.Workbook()
    tabla = wb.active
    tabla.title = "Protocolo"
    tabla.append(["Seccion", "Presion", "Tiempo"])
    for c in tabla[1]:
        c.font = Font(bold=True)
    for i in range(1, 5):
        tabla.append([i, f"{3 + 2 * i} p.s.i", "1 hora"])
    mural = wb.create_sheet("Mural")
    mural["A1"] = "Contraindicaciones"
    mural["A1"].font = Font(bold=True)
    mural["A1"].fill = PatternFill("solid", fgColor="FFDEEAF6")
    mural["A2"] = "Neumotorax no tratado\nFiebre"
    mural["A2"].comment = Comment("Revisar antes de cada sesión", "Docente")
    mural.add_image(XlImagen(str(png)), "D1")
    wb.save(ruta)


def _correr(*args: str, datos: Path) -> subprocess.CompletedProcess:
    entorno = {**os.environ, "BIBLIOTECA_DATOS": str(datos)}
    return subprocess.run([sys.executable, *args], cwd=RAIZ, env=entorno, capture_output=True, text=True, timeout=120)


def test_preparar_cargar_y_auditar(con, datos: Path, tmp_path: Path) -> None:
    mid = bib.crear_materia(con, "Oxigenoterapia", "", "azul", "wind")
    bib.crear_componente(con, mid, None, "HBO", "", ["notas"])
    png = tmp_path / "lamina.png"
    Image.new("RGB", (60, 40), "white").save(png)
    libro = tmp_path / "Oxigeno.xlsx"
    _libro(libro, png)

    r = _correr(str(CARGA / "preparar_fuentes.py"), str(libro), "--materia", "Oxigenoterapia", "--componente", "HBO", datos=datos)
    assert r.returncode == 0, r.stderr
    base = datos / "fuentes" / "oxigenoterapia"
    carpeta = base / "hbo"
    assert (base / "_originales" / "Oxigeno.xlsx").read_bytes() == libro.read_bytes()   # original intacto
    assert openpyxl.load_workbook(carpeta / "hojas" / "Protocolo.xlsx").sheetnames == ["Protocolo"]  # una hoja por archivo
    assert openpyxl.load_workbook(carpeta / "hojas" / "Mural.xlsx")["Mural"]["A2"].comment is not None
    assert (carpeta / "imagenes" / "Mural - D1.png").exists()                           # imagen por su celda ancla
    assert "Revisar antes de cada sesión" in (carpeta / "inspeccion-Mural.txt").read_text()

    (carpeta / "notas").mkdir()
    (carpeta / "notas" / "Contraindicaciones - HBO.md").write_text(
        "## Contraindicaciones\n- Neumotórax no tratado\n- Fiebre\n\n> Comentario de Docente: Revisar antes de cada sesión\n")
    plan = {"materia": "oxigenoterapia", "componente": "HBO", "fuentes": ["hojas/Protocolo.xlsx", "hojas/Mural.xlsx"],
            "correcciones": {"Presion": "Presión"},
            "notas": [{"titulo": "Contraindicaciones - HBO", "archivo": "notas/Contraindicaciones - HBO.md"}],
            "tablas": [{"archivo": "hojas/Protocolo.xlsx", "nombre": "Protocolo", "columna_titulo": "Seccion"}],
            "imagenes": [{"archivo": "imagenes/Mural - D1.png", "titulo": "Lámina", "descripcion": ""}]}
    (carpeta / "plan.json").write_text(json.dumps(plan, ensure_ascii=False))

    r = _correr(str(CARGA / "cargar.py"), str(carpeta / "plan.json"), datos=datos)
    assert r.returncode == 0, r.stdout + r.stderr
    assert list((datos / "respaldos").glob("antes-de-cargar-*.zip"))                 # respaldo previo
    assert list((carpeta / "cargas").glob("*.json"))                                 # registro de lo creado
    r = _correr(str(CARGA / "auditar.py"), str(carpeta / "plan.json"), datos=datos)
    assert r.returncode == 0, r.stdout

    # Una palabra de la fuente que no quedó en la biblioteca se detecta
    con.execute("UPDATE notas SET contenido = replace(contenido, 'Fiebre', '') WHERE titulo = 'Contraindicaciones - HBO'")
    r = _correr(str(CARGA / "auditar.py"), str(carpeta / "plan.json"), datos=datos)
    assert r.returncode == 1 and "fiebre" in r.stdout

    # Volver a cargar sin --reemplazar no duplica nada
    r = _correr(str(CARGA / "cargar.py"), str(carpeta / "plan.json"), datos=datos)
    assert r.returncode != 0 and "Ya existe" in (r.stdout + r.stderr)
