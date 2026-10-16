"""Auditoría de una carga: compara las fuentes, el plan y lo que quedó en la biblioteca.

    .venv/bin/python tools/carga/auditar.py <plan.json> [--ocr]

Revisa:
  1. Cada nota, imagen, tabla y PDF del plan existe en el componente con el contenido exacto del plan
     (texto de la nota, descripción y archivo idéntico de la imagen, filas y celdas de la tabla).
  2. Cobertura: cada palabra de cada celda y comentario de las fuentes aparece en lo cargado. Solo se
     aceptan diferencias declaradas en "correcciones" (ortografía) u "omitidos" (con motivo).
  3. Toda imagen de la carpeta imagenes/ del plan está cargada o declarada en "omitidos".
  4. Con --ocr: palabras que el OCR lee en cada imagen y no están en su descripción (aviso, no error).
Sale con código 1 si hay errores.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import comun  # noqa: F401
import openpyxl
from openpyxl.utils import range_boundaries

from app import db
from app.services import almacenamiento as alm
from app.services import hojas as lector
from app.services.texto import normalizar
from comun import palabras, texto_comentario
from resolver import componente_del_plan


def _celdas_omitidas(plan: dict) -> set[str]:
    salida = set()
    for o in plan.get("omitidos", []):
        if "celda" in o:
            hoja, _, rango = o["celda"].rpartition("!")
            c0, f0, c1, f1 = range_boundaries(rango if ":" in rango else f"{rango}:{rango}")
            for r in range(f0, f1 + 1):
                for c in range(c0, c1 + 1):
                    salida.add(f"{hoja}!{openpyxl.utils.get_column_letter(c)}{r}")
    return salida


def auditar(plan_ruta: Path, con_ocr: bool) -> tuple[list[str], list[str]]:
    plan = json.loads(plan_ruta.read_text(encoding="utf-8"))
    base = plan_ruta.parent
    errores, avisos = [], []
    con = db.conectar()
    cid = componente_del_plan(con, plan)
    if cid is None:
        return [f"No existe el componente «{plan['componente']}» en «{plan['materia']}»."], []

    # 1. Elementos del plan, con su contenido exacto
    for n in plan.get("notas", []):
        esperado = ((base / n["archivo"]).read_text(encoding="utf-8") if "archivo" in n else n["contenido"]).strip() + "\n"
        f = con.execute("SELECT contenido FROM notas WHERE componente_id=? AND titulo=? AND papelera_id IS NULL",
                        (cid, n["titulo"])).fetchone()
        if not f:
            errores.append(f"Falta la nota «{n['titulo']}».")
        elif f["contenido"] != esperado:
            errores.append(f"La nota «{n['titulo']}» no tiene el contenido del plan.")
    for i in plan.get("imagenes", []):
        f = con.execute("SELECT descripcion, hash FROM imagenes WHERE componente_id=? AND titulo=? AND papelera_id IS NULL",
                        (cid, i["titulo"])).fetchone()
        if not f:
            errores.append(f"Falta la imagen «{i['titulo']}».")
            continue
        if f["descripcion"] != (i.get("descripcion", "") or "").strip():
            errores.append(f"La imagen «{i['titulo']}» no tiene la descripción del plan.")
        if f["hash"] != alm.hash_archivo(base / i["archivo"]):
            errores.append(f"El archivo de la imagen «{i['titulo']}» no es el de las fuentes ({i['archivo']}).")
    for t in plan.get("tablas", []):
        f = con.execute("SELECT id FROM tablas WHERE componente_id=? AND nombre=? AND papelera_id IS NULL",
                        (cid, t["nombre"])).fetchone()
        if not f:
            errores.append(f"Falta la tabla «{t['nombre']}».")
            continue
        filas = [json.loads(r[0]) for r in con.execute("SELECT valores FROM filas WHERE tabla_id=? AND huerfana=0", (f["id"],))]
        valores = {normalizar(v) for fila in filas for v in fila.values() if v}
        hoja = lector.leer_archivo(base / t["archivo"], Path(t["archivo"]).name)[0]
        enc = lector.detectar_encabezado(hoja) if "fila_encabezado" not in t else t["fila_encabezado"]
        datos = [r for r in hoja.filas[enc:] if any(str(x).strip() for x in r)]
        if len(datos) != len(filas):
            errores.append(f"La tabla «{t['nombre']}» tiene {len(filas)} filas y el Excel {len(datos)}.")
        faltan = {str(x) for r in datos for x in r if str(x).strip() and normalizar(str(x)) not in valores}
        if faltan:
            errores.append(f"La tabla «{t['nombre']}» no tiene estos valores del Excel: {sorted(faltan)[:10]}")
    for p in plan.get("pdfs", []):
        f = con.execute("SELECT hash FROM pdfs WHERE componente_id=? AND titulo=? AND papelera_id IS NULL",
                        (cid, p["titulo"])).fetchone()
        if not f:
            errores.append(f"Falta el PDF «{p['titulo']}».")
        elif f["hash"] != alm.hash_archivo(base / p["archivo"]):
            errores.append(f"El PDF «{p['titulo']}» no es el archivo de las fuentes.")

    # 2. Cobertura de las fuentes: todas sus palabras están en lo cargado
    corpus: set[str] = set()
    for r in con.execute("SELECT titulo, contenido FROM notas WHERE componente_id=? AND papelera_id IS NULL", (cid,)):
        corpus |= palabras(f"{r[0]} {r[1]}")
    for r in con.execute("SELECT titulo, descripcion FROM imagenes WHERE componente_id=? AND papelera_id IS NULL", (cid,)):
        corpus |= palabras(f"{r[0]} {r[1]}")
    for r in con.execute("""SELECT t.nombre, f.valores FROM filas f JOIN tablas t ON t.id = f.tabla_id
                            WHERE t.componente_id=? AND t.papelera_id IS NULL""", (cid,)):
        corpus |= palabras(r[0] + " " + " ".join(json.loads(r[1]).values()))
    for r in con.execute("SELECT columnas FROM tablas WHERE componente_id=? AND papelera_id IS NULL", (cid,)):
        corpus |= palabras(" ".join(c["nombre"] for c in json.loads(r[0])))   # encabezados de columna
    aceptadas = set()
    for original, corregido in plan.get("correcciones", {}).items():
        if palabras(corregido) <= corpus:
            aceptadas |= palabras(original)
        else:
            errores.append(f"La corrección «{original}» → «{corregido}» no aparece en lo cargado.")
    omitidas = _celdas_omitidas(plan)
    revisadas = 0
    for fuente in plan.get("fuentes", []):
        ruta = base / fuente
        if ruta.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        for ws in openpyxl.load_workbook(ruta, data_only=True).worksheets:   # el valor que se ve, como la app
            for fila in ws.iter_rows():
                for c in fila:
                    textos = [lector.texto_celda(c.value, c.number_format)] if c.value is not None else []   # como lo muestra la app
                    if c.comment:
                        textos.append(texto_comentario(c.comment))
                    for tx in textos:
                        revisadas += 1
                        if f"{ws.title}!{c.coordinate}" in omitidas:
                            continue
                        faltan = palabras(tx) - corpus - aceptadas
                        if faltan:
                            errores.append(f"{ws.title}!{c.coordinate}: no está en lo cargado {sorted(faltan)} — «{tx[:80]}»")
    avisos.append(f"Celdas y comentarios revisados: {revisadas}")

    # 3. Imágenes de las fuentes sin cargar
    en_plan = {i["archivo"] for i in plan.get("imagenes", [])} | {o["archivo"] for o in plan.get("omitidos", []) if "archivo" in o}
    usadas_como_fuente = {i.get("fuente") for i in plan.get("imagenes", []) if i.get("fuente")}
    carpeta = base / "imagenes"
    for archivo in sorted(carpeta.glob("*")) if carpeta.exists() else []:
        rel = f"imagenes/{archivo.name}"
        if archivo.is_file() and rel not in en_plan and rel not in usadas_como_fuente:
            errores.append(f"La imagen de las fuentes «{rel}» no está cargada ni declarada en «omitidos».")

    # 4. OCR (opcional): palabras de la imagen que no están en su descripción
    if con_ocr:
        from PIL import Image

        from app.services import ocr
        for i in plan.get("imagenes", []):
            leido = palabras(ocr.reconocer(Image.open(base / i.get("fuente", i["archivo"])).convert("RGB")))
            propio = palabras(f"{i['titulo']} {i.get('descripcion', '')}") | aceptadas
            dif = sorted(p for p in leido - propio if len(p) >= 4 and not p.isdigit())
            if dif:
                avisos.append(f"OCR «{i['titulo']}»: lee {dif} que no está en la descripción (revisar a ojo).")
    return errores, avisos


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("plan", type=Path)
    ap.add_argument("--ocr", action="store_true")
    a = ap.parse_args()
    errores, avisos = auditar(a.plan, a.ocr)
    for x in avisos:
        print(f"AVISO  {x}")
    for x in errores:
        print(f"ERROR  {x}")
    print(f"\n{'CORRECTO' if not errores else f'{len(errores)} ERRORES'}")
    sys.exit(1 if errores else 0)
