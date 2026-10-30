"""Guarda los archivos entregados en data/fuentes/ y los separa para cargarlos.

    .venv/bin/python tools/carga/preparar_fuentes.py <archivo>... --materia "Oxigenoterapia" \
        (--componente "HBO" | --hoja "HBO=HBO" --hoja "Oxivenacion=Oxivenación" ...)

Estructura que deja (nada se modifica después de copiado):
    data/fuentes/<materia>/_originales/<archivo tal cual llegó>
    data/fuentes/<materia>/<componente>/hojas/<hoja>.xlsx      una hoja por archivo (valores, formato, comentarios)
    data/fuentes/<materia>/<componente>/imagenes/<hoja> - <celda>.<ext>   imágenes del Excel, por su celda ancla
    data/fuentes/<materia>/<componente>/inspeccion-<hoja>.txt|.json      descripción de bloques (ver inspeccionar.py)
    data/fuentes/<materia>/<componente>/pdf|imagenes/...      PDF e imágenes entregados sueltos
    data/fuentes/<materia>/<componente>/documentos/<archivo>.docx   documentos de Word (se pasan a notas)

--componente manda todas las hojas a un componente; --hoja "Hoja=Componente" reparte por hoja
(las hojas sin asignar se informan y no se copian).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

import comun  # noqa: F401
import openpyxl

from app.config import rutas
from app.services.almacenamiento import sanear_nombre_archivo, slugify
from inspeccionar import imagenes_por_hoja, inspeccionar, resumen

EXT_IMAGEN = {".png", ".jpg", ".jpeg", ".webp", ".gif"}


def _hash(ruta: Path) -> str:
    return hashlib.sha256(ruta.read_bytes()).hexdigest()


def _copiar_sin_pisar(origen: Path, destino: Path) -> Path:
    """Copia; si ya existe un archivo distinto con ese nombre, agrega la fecha al nombre (nunca reemplaza)."""
    destino.parent.mkdir(parents=True, exist_ok=True)
    if destino.exists():
        if _hash(destino) == _hash(origen):
            return destino
        destino = destino.with_name(f"{destino.stem} ({datetime.now():%Y-%m-%d %H-%M}){destino.suffix}")
    shutil.copy2(origen, destino)
    return destino


def separar_hoja(archivo: Path, hoja: str, destino: Path) -> None:
    """Deja solo `hoja` en una copia del libro (conserva valores, formato, combinadas, comentarios e imágenes).
    Las fórmulas se guardan con su resultado (el valor que se ve): openpyxl no recalcula, y una fórmula guardada
    sin su valor quedaría vacía para la app. El original, con sus fórmulas, queda en _originales/."""
    wb = openpyxl.load_workbook(archivo, data_only=True)
    for ws in list(wb.worksheets):
        if ws.title != hoja:
            wb.remove(ws)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    # Comprobación: las celdas de la copia son idénticas a las del original
    a = openpyxl.load_workbook(archivo, data_only=True)[hoja]
    b = openpyxl.load_workbook(destino, data_only=True)[hoja]
    def valor(v):  # las fórmulas de matriz son objetos: se comparan por su rango y su texto
        return (getattr(v, "ref", None), v.text) if hasattr(v, "text") else v

    va = {c.coordinate: valor(c.value) for f in a.iter_rows() for c in f if c.value is not None}
    vb = {c.coordinate: valor(c.value) for f in b.iter_rows() for c in f if c.value is not None}
    if va != vb:
        raise SystemExit(f"La copia de la hoja «{hoja}» no coincide con el original; no se continúa.")


def extraer_imagenes(archivo: Path, hoja: str, carpeta: Path) -> list[Path]:
    salida = []
    with zipfile.ZipFile(archivo) as z:
        for img in imagenes_por_hoja(archivo).get(hoja, []):
            ext = Path(img["media"]).suffix.lower()
            nombre = sanear_nombre_archivo(f"{hoja} - {img['desde'] or 'sin celda'}{ext}")
            destino = carpeta / nombre
            n = 2
            while destino.exists() and destino.read_bytes() != z.read(img["media"]):
                destino = carpeta / sanear_nombre_archivo(f"{hoja} - {img['desde']} ({n}){ext}")
                n += 1
            carpeta.mkdir(parents=True, exist_ok=True)
            destino.write_bytes(z.read(img["media"]))
            salida.append(destino)
    return salida


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("archivos", nargs="+", type=Path)
    ap.add_argument("--materia", required=True)
    ap.add_argument("--componente", help="componente para todas las hojas / archivos")
    ap.add_argument("--hoja", action="append", default=[], help='"Nombre de hoja=Componente"')
    a = ap.parse_args()
    por_hoja = dict(h.split("=", 1) for h in a.hoja)
    if not a.componente and not por_hoja:
        ap.error("indica --componente o al menos un --hoja")
    base = rutas.fuentes / slugify(a.materia)
    informe = {"materia": a.materia, "fecha": datetime.now().isoformat(timespec="seconds"), "archivos": []}
    for archivo in a.archivos:
        original = _copiar_sin_pisar(archivo, base / "_originales" / archivo.name)
        registro = {"original": rutas.relativa(original), "hojas": [], "sin_asignar": []}
        ext = archivo.suffix.lower()
        if ext in {".xlsx", ".xlsm"}:
            for hoja in openpyxl.load_workbook(original, read_only=True).sheetnames:
                componente = por_hoja.get(hoja, a.componente)
                if not componente:
                    registro["sin_asignar"].append(hoja)
                    continue
                carpeta = base / slugify(componente)
                destino = carpeta / "hojas" / sanear_nombre_archivo(f"{hoja}.xlsx")
                separar_hoja(original, hoja, destino)
                imgs = extraer_imagenes(original, hoja, carpeta / "imagenes")
                info = inspeccionar(destino)
                info["hojas"][0]["imagenes_extraidas"] = [rutas.relativa(p) for p in imgs]
                (carpeta / f"inspeccion-{destino.stem}.json").write_text(json.dumps(info, ensure_ascii=False, indent=2), "utf-8")
                (carpeta / f"inspeccion-{destino.stem}.txt").write_text(resumen(info), "utf-8")
                registro["hojas"].append({"hoja": hoja, "componente": componente, "archivo": rutas.relativa(destino),
                                          "imagenes": [rutas.relativa(p) for p in imgs]})
        elif ext in EXT_IMAGEN | {".pdf", ".csv", ".docx"}:
            if not a.componente:
                ap.error(f"{archivo.name}: indica --componente para archivos que no son Excel")
            sub = {".pdf": "pdf", ".csv": "hojas", ".docx": "documentos"}.get(ext, "imagenes")
            copia = _copiar_sin_pisar(original, base / slugify(a.componente) / sub / archivo.name)
            registro["hojas"].append({"componente": a.componente, "archivo": rutas.relativa(copia)})
        else:
            registro["sin_asignar"].append(f"formato no reconocido: {archivo.name}")
        informe["archivos"].append(registro)
    base.mkdir(parents=True, exist_ok=True)
    (base / f"preparacion-{datetime.now():%Y-%m-%d_%H-%M-%S}.json").write_text(
        json.dumps(informe, ensure_ascii=False, indent=2), "utf-8")
    print(json.dumps(informe, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
