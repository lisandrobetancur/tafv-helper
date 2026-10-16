"""Copia un rango de una hoja (una tabla dentro de una hoja con recuadros) a un Excel propio, tal cual.

    .venv/bin/python tools/carga/extraer_rango.py <hoja.xlsx> --rango Q2:S12 --salida <carpeta>/tablas/<Nombre>.xlsx

Se copian valores, negrita, rellenos y celdas combinadas del rango, empezando en A1. No se completa,
ordena ni corrige nada: si la tabla necesita correcciones de ortografía se hacen en el plan de carga.
"""
from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

import comun  # noqa: F401
import openpyxl
from openpyxl.utils import range_boundaries


def extraer(origen: Path, rango: str, salida: Path, hoja: str | None = None) -> int:
    wb = openpyxl.load_workbook(origen, data_only=True)
    ws = wb[hoja] if hoja else wb.worksheets[0]
    c0, f0, c1, f1 = range_boundaries(rango)
    nuevo = openpyxl.Workbook()
    dst = nuevo.active
    dst.title = (hoja or ws.title)[:31]
    for r in range(f0, f1 + 1):
        for c in range(c0, c1 + 1):
            o = ws.cell(r, c)
            d = dst.cell(r - f0 + 1, c - c0 + 1, o.value)
            if o.has_style:
                d.font, d.fill, d.alignment, d.number_format = copy(o.font), copy(o.fill), copy(o.alignment), o.number_format
    for m in ws.merged_cells.ranges:
        if m.min_col >= c0 and m.max_col <= c1 and m.min_row >= f0 and m.max_row <= f1:
            dst.merge_cells(start_row=m.min_row - f0 + 1, start_column=m.min_col - c0 + 1,
                            end_row=m.max_row - f0 + 1, end_column=m.max_col - c0 + 1)
    salida.parent.mkdir(parents=True, exist_ok=True)
    nuevo.save(salida)
    return f1 - f0 + 1


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("origen", type=Path)
    ap.add_argument("--rango", required=True)
    ap.add_argument("--salida", required=True, type=Path)
    ap.add_argument("--hoja")
    a = ap.parse_args()
    print(f"{extraer(a.origen, a.rango, a.salida, a.hoja)} filas copiadas a {a.salida}")
