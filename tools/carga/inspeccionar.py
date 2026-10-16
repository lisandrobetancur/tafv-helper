"""Describe una hoja de Excel para decidir cómo cargarla: bloques de celdas, encabezados, comentarios e imágenes.

    .venv/bin/python tools/carga/inspeccionar.py <archivo.xlsx> [--json salida.json]

Cada «bloque» es un grupo de celdas con contenido que se tocan (las combinadas cuentan como una).
La sugerencia «tabla» (filas × columnas completas con encabezado) o «recuadros» es solo una pista:
quien carga decide mirando el texto de cada bloque.
"""
from __future__ import annotations

import argparse
import json
import posixpath
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import comun  # noqa: F401  (agrega la raíz del proyecto al path)
import openpyxl
from openpyxl.utils import get_column_letter

from comun import autor_comentario, texto_comentario


def _relleno(celda) -> str | None:
    f = celda.fill
    if f is None or f.fill_type is None:
        return None
    rgb = f.fgColor.rgb if f.fgColor is not None and f.fgColor.type == "rgb" else None
    return rgb if isinstance(rgb, str) and rgb not in ("00000000", "FFFFFFFF") else None


NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}


def imagenes_por_hoja(archivo: Path) -> dict[str, list[dict]]:
    """{nombre de hoja: [{"media": "xl/media/image1.png", "desde": "A36", "hasta": "H50"}]} leyendo el XML del libro
    (sirve para archivos de Excel, Google Sheets o LibreOffice: rutas relativas o absolutas, con o sin prefijos)."""
    salida: dict[str, list[dict]] = {}
    with zipfile.ZipFile(archivo) as z:
        nombres = set(z.namelist())

        def rels(ruta: str) -> dict[str, str]:
            rel = posixpath.join(posixpath.dirname(ruta), "_rels", posixpath.basename(ruta) + ".rels")
            if rel not in nombres:
                return {}
            mapa = {}
            for r in ET.fromstring(z.read(rel)).findall("rel:Relationship", NS):
                destino = r.get("Target", "")
                destino = destino.lstrip("/") if destino.startswith("/") else posixpath.join(posixpath.dirname(ruta), destino)
                mapa[r.get("Id")] = posixpath.normpath(destino)
            return mapa

        rels_libro = rels("xl/workbook.xml")
        for hoja in ET.fromstring(z.read("xl/workbook.xml")).iter(f"{{{NS['m']}}}sheet"):
            nombre = hoja.get("name", "")
            ruta_hoja = rels_libro.get(hoja.get(f"{{{NS['r']}}}id"))
            lista = salida.setdefault(nombre, [])
            for destino in rels(ruta_hoja).values() if ruta_hoja else []:
                if "/drawings/" not in destino or not destino.endswith(".xml") or destino not in nombres:
                    continue
                medios = rels(destino)
                raiz = ET.fromstring(z.read(destino))
                for ancla in [*raiz.findall("xdr:twoCellAnchor", NS), *raiz.findall("xdr:oneCellAnchor", NS),
                              *raiz.findall("xdr:absoluteAnchor", NS)]:
                    blip = ancla.find(".//a:blip", NS)
                    emb = blip.get(f"{{{NS['r']}}}embed") if blip is not None else None
                    if emb not in medios:
                        continue

                    def celda(etiqueta: str) -> str | None:
                        p = ancla.find(f"xdr:{etiqueta}", NS)
                        if p is None:
                            return None
                        col, fila = p.find("xdr:col", NS), p.find("xdr:row", NS)
                        return f"{get_column_letter(int(col.text) + 1)}{int(fila.text) + 1}"

                    lista.append({"media": medios[emb], "desde": celda("from"), "hasta": celda("to")})
    return salida


def bloques(ws) -> list[dict]:
    """Agrupa las celdas con contenido en bloques conectados (vecinas arriba/abajo/izquierda/derecha)."""
    ocupadas: dict[tuple[int, int], tuple[int, int]] = {}   # celda -> celda ancla (la de arriba a la izquierda)
    combinadas = {}
    for rango in ws.merged_cells.ranges:
        ancla = (rango.min_row, rango.min_col)
        combinadas[ancla] = rango.coord
        if ws.cell(*ancla).value is not None:
            for r in range(rango.min_row, rango.max_row + 1):
                for c in range(rango.min_col, rango.max_col + 1):
                    ocupadas[(r, c)] = ancla
    for fila in ws.iter_rows():
        for c in fila:
            if c.value is not None and str(c.value).strip() != "":
                ocupadas.setdefault((c.row, c.column), (c.row, c.column))
    vistos, salida = set(), []
    for inicio in sorted(ocupadas):
        if inicio in vistos:
            continue
        pila, grupo = [inicio], []
        vistos.add(inicio)
        while pila:
            r, c = pila.pop()
            grupo.append((r, c))
            for v in ((r + 1, c), (r - 1, c), (r, c + 1), (r, c - 1)):
                if v in ocupadas and v not in vistos:
                    vistos.add(v)
                    pila.append(v)
        anclas = sorted({ocupadas[p] for p in grupo})
        f0, f1 = min(p[0] for p in grupo), max(p[0] for p in grupo)
        c0, c1 = min(p[1] for p in grupo), max(p[1] for p in grupo)
        celdas = []
        for r, c in anclas:
            x = ws.cell(r, c)
            celdas.append({"celda": x.coordinate, "combinada": combinadas.get((r, c)), "texto": str(x.value),
                           "negrita": bool(x.font and x.font.b), "relleno": _relleno(x)})
        filas = f1 - f0 + 1
        columnas = c1 - c0 + 1
        densidad = len(anclas) / (filas * columnas)
        primera = [x for x in celdas if openpyxl.utils.cell.coordinate_from_string(x["celda"])[1] == f0]
        encabezado = len(primera) >= 2 and all(x["negrita"] or x["relleno"] for x in primera)
        con_combinadas = sum(1 for x in celdas if x["combinada"])
        sugerencia = "tabla" if (columnas >= 2 and filas >= 3 and densidad >= 0.7 and encabezado
                                 and con_combinadas <= max(1, len(primera) // 2)) else "recuadros"
        salida.append({"rango": f"{get_column_letter(c0)}{f0}:{get_column_letter(c1)}{f1}", "filas": filas,
                       "columnas": columnas, "densidad": round(densidad, 2), "encabezado_en_primera_fila": encabezado,
                       "sugerencia": sugerencia, "celdas": celdas})
    return salida


def inspeccionar(archivo: Path) -> dict:
    wb = openpyxl.load_workbook(archivo)
    imagenes = imagenes_por_hoja(archivo)
    hojas = []
    for ws in wb.worksheets:
        comentarios = [{"celda": c.coordinate, "autor": autor_comentario(c.comment), "texto": texto_comentario(c.comment)}
                       for fila in ws.iter_rows() for c in fila if c.comment]
        hojas.append({"hoja": ws.title, "dimension": ws.dimensions, "bloques": bloques(ws),
                      "comentarios": comentarios, "imagenes": imagenes.get(ws.title, [])})
    return {"archivo": str(archivo), "hojas": hojas}


def resumen(info: dict) -> str:
    lineas = [f"Archivo: {info['archivo']}"]
    for h in info["hojas"]:
        lineas.append(f"\n=== Hoja «{h['hoja']}» ({h['dimension']}) — {len(h['bloques'])} bloques, "
                      f"{len(h['comentarios'])} comentarios, {len(h['imagenes'])} imágenes")
        for b in h["bloques"]:
            lineas.append(f"\n--- Bloque {b['rango']} · {b['filas']}×{b['columnas']} · densidad {b['densidad']} · "
                          f"sugerencia: {b['sugerencia'].upper()}")
            for c in b["celdas"]:
                marcas = "".join(["N" if c["negrita"] else "", f" {c['relleno']}" if c["relleno"] else "",
                                  f" combinada {c['combinada']}" if c["combinada"] else ""])
                lineas.append(f"  {c['celda']} [{marcas.strip()}] {c['texto']!r}")
        for c in h["comentarios"]:
            lineas.append(f"  COMENTARIO {c['celda']} ({c['autor']}): {c['texto']!r}")
        for i in h["imagenes"]:
            lineas.append(f"  IMAGEN {i['media']} anclada en {i['desde']}" + (f":{i['hasta']}" if i["hasta"] else ""))
    return "\n".join(lineas)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("archivo", type=Path)
    ap.add_argument("--json", type=Path, help="guardar también la descripción completa en JSON")
    a = ap.parse_args()
    datos = inspeccionar(a.archivo)
    if a.json:
        a.json.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    print(resumen(datos))
