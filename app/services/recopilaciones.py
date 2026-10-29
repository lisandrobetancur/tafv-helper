"""Recopilaciones: fragmentos (fichas, campos, textos subrayados, imágenes y recortes) que se van juntando
desde cualquier materia, se ordenan y se exportan a PDF o Word.

Cada fragmento es una copia fija tomada al agregarlo: si después cambia o se borra el material de la
biblioteca, la recopilación no cambia. Viven en su propia base de datos (`data/recopilaciones/`), que no
va en Git: son de trabajo personal y se borran cuando la usuaria quiera.
"""
from __future__ import annotations

import html
import io
import json
import re
import sqlite3
import uuid
from collections.abc import Iterator
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from PIL import Image

from ..config import rutas
from ..errores import ErrorUsuario, NoEncontrado
from . import elementos
from . import imagenes as img_srv
from . import imagenes_fila as imf
from . import notas as notas_srv
from . import pdfs as pdf_srv
from . import tablas as tab

TIPOS = {"campos", "texto", "imagen", "seccion"}
TAMANOS = {"pequena": ("Pequeña", 170), "mediana": ("Mediana", 300), "grande": ("Grande", 504)}  # ancho en puntos
ANCHO_PAGINA, ALTO_MAXIMO = 504, 600      # carta con márgenes de 54 pt
DPI_PDF = 200                             # resolución al copiar páginas o recortes de un PDF
LADO_MAXIMO = 2400                        # las imágenes copiadas se reducen hasta este lado (px)
TITULO_POR_DEFECTO = "Mi recopilación"

ESQUEMA = """
CREATE TABLE IF NOT EXISTS recopilaciones (
    id           INTEGER PRIMARY KEY,
    titulo       TEXT NOT NULL,
    descripcion  TEXT NOT NULL DEFAULT '',
    creado       TEXT NOT NULL,
    actualizado  TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS fragmentos (
    id              INTEGER PRIMARY KEY,
    recopilacion_id INTEGER NOT NULL REFERENCES recopilaciones(id) ON DELETE CASCADE,
    orden           INTEGER NOT NULL,
    tipo            TEXT NOT NULL,            -- campos | texto | imagen | seccion
    titulo          TEXT NOT NULL DEFAULT '',
    contenido       TEXT NOT NULL DEFAULT '', -- texto del fragmento
    formato         TEXT NOT NULL DEFAULT 'texto',  -- texto | markdown
    campos          TEXT NOT NULL DEFAULT '[]',     -- JSON [[columna, valor], …] (tipo campos)
    archivo         TEXT,                     -- imagen copiada, en data/recopilaciones/imagenes/
    ancho           INTEGER,
    alto            INTEGER,
    tamano          TEXT NOT NULL DEFAULT 'mediana',
    comentario      TEXT NOT NULL DEFAULT '', -- lo que agrega la usuaria (Markdown)
    origen          TEXT NOT NULL DEFAULT '', -- «Materia › Componente › …»
    origen_url      TEXT NOT NULL DEFAULT '',
    creado          TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS ix_fragmentos ON fragmentos(recopilacion_id, orden);
CREATE TABLE IF NOT EXISTS ajustes (clave TEXT PRIMARY KEY, valor TEXT NOT NULL);
"""


def carpeta() -> Path:
    return rutas.datos / "recopilaciones"


def carpeta_imagenes() -> Path:
    return carpeta() / "imagenes"


def ahora() -> str:
    return datetime.now().isoformat(timespec="seconds")


def conectar() -> sqlite3.Connection:
    carpeta_imagenes().mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(carpeta() / "recopilaciones.db", timeout=10, check_same_thread=False, isolation_level=None)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    con.execute("PRAGMA journal_mode = DELETE")
    con.executescript(ESQUEMA)
    return con


def obtener_rdb() -> Iterator[sqlite3.Connection]:
    """Dependencia de FastAPI: una conexión a las recopilaciones por petición."""
    con = conectar()
    try:
        yield con
    finally:
        con.close()


# ------------------------------------------------------------------ modelos
@dataclass
class Fragmento:
    id: int
    recopilacion_id: int
    orden: int
    tipo: str
    titulo: str
    contenido: str
    formato: str
    campos: list[list[str]]
    archivo: str | None
    ancho: int | None
    alto: int | None
    tamano: str
    comentario: str
    origen: str
    origen_url: str

    @classmethod
    def desde_fila(cls, f: sqlite3.Row) -> Fragmento:
        return cls(id=f["id"], recopilacion_id=f["recopilacion_id"], orden=f["orden"], tipo=f["tipo"], titulo=f["titulo"],
                   contenido=f["contenido"], formato=f["formato"], campos=json.loads(f["campos"]), archivo=f["archivo"],
                   ancho=f["ancho"], alto=f["alto"], tamano=f["tamano"], comentario=f["comentario"], origen=f["origen"],
                   origen_url=f["origen_url"])

    @property
    def html(self) -> str:
        return notas_srv.renderizar(self.contenido, self.formato) if self.contenido.strip() else ""

    @property
    def html_comentario(self) -> str:
        return notas_srv.renderizar(self.comentario, "markdown") if self.comentario.strip() else ""

    def medidas(self) -> tuple[int, int]:
        """Ancho y alto (en puntos) de la imagen en el documento: el tamaño elegido, sin agrandar
        recortes pequeños (se verían borrosos) y sin pasar del alto de una página."""
        if not self.ancho or not self.alto:
            return TAMANOS["mediana"][1], TAMANOS["mediana"][1]
        ancho = min(TAMANOS.get(self.tamano, TAMANOS["mediana"])[1], ANCHO_PAGINA, max(self.ancho * 72 // 110, 60))
        alto = ancho * self.alto / self.ancho
        if alto > ALTO_MAXIMO:
            ancho, alto = ancho * ALTO_MAXIMO / alto, ALTO_MAXIMO
        return int(ancho), int(alto)


@dataclass
class Recopilacion:
    id: int
    titulo: str
    descripcion: str
    creado: str
    actualizado: str
    num_fragmentos: int = 0
    fragmentos: list[Fragmento] = field(default_factory=list)


def _recopilacion(f: sqlite3.Row) -> Recopilacion:
    return Recopilacion(id=f["id"], titulo=f["titulo"], descripcion=f["descripcion"], creado=f["creado"],
                        actualizado=f["actualizado"], num_fragmentos=f["n"])


_SELECT = "SELECT r.*, (SELECT COUNT(*) FROM fragmentos g WHERE g.recopilacion_id = r.id) AS n FROM recopilaciones r"


# ------------------------------------------------------------ recopilaciones
def listar(rcon: sqlite3.Connection) -> list[Recopilacion]:
    return [_recopilacion(f) for f in rcon.execute(f"{_SELECT} ORDER BY r.actualizado DESC, r.id DESC")]


def obtener(rcon: sqlite3.Connection, rid: int, con_fragmentos: bool = True) -> Recopilacion:
    f = rcon.execute(f"{_SELECT} WHERE r.id = ?", (rid,)).fetchone()
    if f is None:
        raise NoEncontrado("La recopilación")
    r = _recopilacion(f)
    if con_fragmentos:
        r.fragmentos = [Fragmento.desde_fila(g) for g in
                        rcon.execute("SELECT * FROM fragmentos WHERE recopilacion_id = ? ORDER BY orden, id", (rid,))]
    return r


def _validar_titulo(titulo: str) -> str:
    titulo = " ".join((titulo or "").split())
    if not titulo:
        raise ErrorUsuario("Escribe un nombre para la recopilación.")
    if len(titulo) > 120:
        raise ErrorUsuario("El nombre es demasiado largo (máximo 120 caracteres).")
    return titulo


def crear(rcon: sqlite3.Connection, titulo: str, descripcion: str = "") -> Recopilacion:
    t = ahora()
    rid = rcon.execute("INSERT INTO recopilaciones (titulo, descripcion, creado, actualizado) VALUES (?,?,?,?)",
                       (_validar_titulo(titulo), descripcion.strip(), t, t)).lastrowid
    activar(rcon, rid)
    return obtener(rcon, rid)


def editar(rcon: sqlite3.Connection, rid: int, titulo: str, descripcion: str) -> None:
    obtener(rcon, rid, False)
    rcon.execute("UPDATE recopilaciones SET titulo=?, descripcion=?, actualizado=? WHERE id=?",
                 (_validar_titulo(titulo), descripcion.strip()[:2000], ahora(), rid))


def eliminar(rcon: sqlite3.Connection, rid: int) -> None:
    """Borra la recopilación y sus imágenes copiadas (no toca el material de la biblioteca)."""
    obtener(rcon, rid, False)
    archivos = [f[0] for f in rcon.execute("SELECT archivo FROM fragmentos WHERE recopilacion_id=? AND archivo IS NOT NULL", (rid,))]
    rcon.execute("DELETE FROM recopilaciones WHERE id=?", (rid,))
    for a in archivos:
        _borrar_imagen(a)
    if activa_id(rcon) == rid:
        rcon.execute("DELETE FROM ajustes WHERE clave='activa'")


def _tocar(rcon: sqlite3.Connection, rid: int) -> None:
    rcon.execute("UPDATE recopilaciones SET actualizado=? WHERE id=?", (ahora(), rid))


def activar(rcon: sqlite3.Connection, rid: int) -> None:
    rcon.execute("INSERT INTO ajustes (clave, valor) VALUES ('activa', ?) ON CONFLICT (clave) DO UPDATE SET valor=excluded.valor",
                 (str(rid),))


def activa_id(rcon: sqlite3.Connection) -> int | None:
    f = rcon.execute("SELECT valor FROM ajustes WHERE clave='activa'").fetchone()
    if f is None:
        return None
    rid = int(f[0])
    return rid if rcon.execute("SELECT 1 FROM recopilaciones WHERE id=?", (rid,)).fetchone() else None


def activa(rcon: sqlite3.Connection, crear_si_falta: bool = False) -> Recopilacion | None:
    """La recopilación donde van los fragmentos nuevos (si no hay ninguna, se crea una al agregar)."""
    rid = activa_id(rcon)
    if rid is None:
        ultima = rcon.execute("SELECT id FROM recopilaciones ORDER BY actualizado DESC, id DESC LIMIT 1").fetchone()
        if ultima:
            rid = ultima[0]
            activar(rcon, rid)
        elif crear_si_falta:
            return crear(rcon, TITULO_POR_DEFECTO)
        else:
            return None
    return obtener(rcon, rid, False)


# ---------------------------------------------------------------- fragmentos
def obtener_fragmento(rcon: sqlite3.Connection, fid: int) -> Fragmento:
    f = rcon.execute("SELECT * FROM fragmentos WHERE id=?", (fid,)).fetchone()
    if f is None:
        raise NoEncontrado("El fragmento")
    return Fragmento.desde_fila(f)


def _insertar(rcon: sqlite3.Connection, rid: int, tipo: str, **datos) -> int:
    orden = rcon.execute("SELECT COALESCE(MAX(orden), 0) + 1 FROM fragmentos WHERE recopilacion_id=?", (rid,)).fetchone()[0]
    if "campos" in datos:
        datos["campos"] = json.dumps(datos["campos"], ensure_ascii=False)
    columnas = ["recopilacion_id", "orden", "tipo", "creado", *datos]
    fid = rcon.execute(f"INSERT INTO fragmentos ({', '.join(columnas)}) VALUES ({', '.join('?' * len(columnas))})",
                       (rid, orden, tipo, ahora(), *datos.values())).lastrowid
    _tocar(rcon, rid)
    return fid


def agregar_seccion(rcon: sqlite3.Connection, rid: int, titulo: str) -> int:
    obtener(rcon, rid, False)
    titulo = " ".join((titulo or "").split())
    if not titulo:
        raise ErrorUsuario("Escribe el título de la sección.")
    return _insertar(rcon, rid, "seccion", titulo=titulo[:200])


def agregar_texto_propio(rcon: sqlite3.Connection, rid: int, titulo: str, contenido: str) -> int:
    obtener(rcon, rid, False)
    if not contenido.strip() and not titulo.strip():
        raise ErrorUsuario("Escribe algo de texto.")
    return _insertar(rcon, rid, "texto", titulo=" ".join(titulo.split())[:200], contenido=contenido.strip(), formato="markdown")


def editar_fragmento(rcon: sqlite3.Connection, fid: int, titulo: str, comentario: str, contenido: str | None = None,
                     tamano: str | None = None) -> Fragmento:
    fr = obtener_fragmento(rcon, fid)
    titulo = " ".join((titulo or "").split())[:200]
    if fr.tipo == "seccion" and not titulo:
        raise ErrorUsuario("La sección necesita un título.")
    cambios = {"titulo": titulo, "comentario": (comentario or "").strip()[:20000]}
    if contenido is not None and fr.tipo == "texto":
        if not contenido.strip():
            raise ErrorUsuario("El texto no puede quedar vacío; si ya no lo necesitas, quita el fragmento.")
        cambios["contenido"] = contenido.strip()
    if tamano is not None and fr.tipo == "imagen":
        cambios["tamano"] = tamano if tamano in TAMANOS else "mediana"
    rcon.execute(f"UPDATE fragmentos SET {', '.join(f'{k}=?' for k in cambios)} WHERE id=?", (*cambios.values(), fid))
    _tocar(rcon, fr.recopilacion_id)
    return obtener_fragmento(rcon, fid)


def eliminar_fragmento(rcon: sqlite3.Connection, fid: int) -> int:
    fr = obtener_fragmento(rcon, fid)
    rcon.execute("DELETE FROM fragmentos WHERE id=?", (fid,))
    if fr.archivo:
        _borrar_imagen(fr.archivo)
    _tocar(rcon, fr.recopilacion_id)
    return fr.recopilacion_id


def mover(rcon: sqlite3.Connection, fid: int, paso: int) -> int:
    """Sube (-1) o baja (+1) un fragmento una posición."""
    fr = obtener_fragmento(rcon, fid)
    ids = [f[0] for f in rcon.execute("SELECT id FROM fragmentos WHERE recopilacion_id=? ORDER BY orden, id", (fr.recopilacion_id,))]
    i = ids.index(fid)
    j = i + (1 if paso > 0 else -1)
    if 0 <= j < len(ids):
        ids[i], ids[j] = ids[j], ids[i]
        ordenar(rcon, fr.recopilacion_id, ids)
    return fr.recopilacion_id


def ordenar(rcon: sqlite3.Connection, rid: int, ids: list[int]) -> None:
    actuales = {f[0] for f in rcon.execute("SELECT id FROM fragmentos WHERE recopilacion_id=?", (rid,))}
    if set(ids) != actuales:
        raise ErrorUsuario("La lista cambió mientras la ordenabas. Recarga la página e inténtalo de nuevo.")
    rcon.execute("BEGIN")
    try:
        for orden, fid in enumerate(ids, start=1):
            rcon.execute("UPDATE fragmentos SET orden=? WHERE id=?", (orden, fid))
        rcon.execute("COMMIT")
    except BaseException:
        rcon.execute("ROLLBACK")
        raise
    _tocar(rcon, rid)


# ----------------------------------------------------------------- captura
def _ubicacion(con: sqlite3.Connection, tipo: str, ref_id: int) -> elementos.Elemento:
    el = elementos.resolver(con, [(tipo, ref_id)])
    if not el:
        raise NoEncontrado("El elemento")
    return el[0]


def _origen_fila(con: sqlite3.Connection, fila_id: int) -> tuple[str, str]:
    fila = tab.obtener_fila(con, fila_id)
    t = tab.obtener(con, fila.tabla_id)
    el = _ubicacion(con, "fila", fila_id)
    return f"{el.ubicacion} › {t.nombre} › {fila.titulo.splitlines()[0][:80]}", el.url


def origen_de(con: sqlite3.Connection, tipo: str, ref_id: int, pagina: int | None = None) -> tuple[str, str]:
    """(texto legible, enlace) de dónde viene un fragmento."""
    if tipo == "fila":
        return _origen_fila(con, ref_id)
    if tipo == "imagen_fila":
        img = imf.obtener(con, ref_id)
        texto, url = _origen_fila(con, img.fila_id)
        return f"{texto} › imagen «{img.titulo}»", url
    if tipo == "pdf":
        el = _ubicacion(con, "pdf", ref_id)
        if pagina:
            return f"{el.ubicacion} › {el.titulo} › página {pagina}", f"{el.url}?pagina={pagina}"
        return f"{el.ubicacion} › {el.titulo}", el.url
    if tipo in ("nota", "imagen"):
        el = _ubicacion(con, tipo, ref_id)
        return f"{el.ubicacion} › {el.titulo}", el.url
    raise ErrorUsuario("No se reconoce de dónde viene este fragmento.")


def agregar_fila(con: sqlite3.Connection, rcon: sqlite3.Connection, fila_id: int, rid: int) -> int:
    """La ficha completa: su título y los campos con datos, en el orden del Excel (sin imágenes)."""
    fila = tab.obtener_fila(con, fila_id)
    t = tab.obtener(con, fila.tabla_id)
    campos = [[c["nombre"], fila.valores.get(c["clave"], "").strip()] for c in t.columnas
              if c["clave"] != t.columna_titulo and fila.valores.get(c["clave"], "").strip()]
    origen, url = origen_de(con, "fila", fila_id)
    return _insertar(rcon, rid, "campos", titulo=fila.titulo.strip()[:300], campos=campos, origen=origen, origen_url=url)


def agregar_campo(con: sqlite3.Connection, rcon: sqlite3.Connection, fila_id: int, clave: str, rid: int) -> int:
    fila = tab.obtener_fila(con, fila_id)
    t = tab.obtener(con, fila.tabla_id)
    nombre = t.nombre_columna(clave)
    valor = (fila.valores.get(clave) or "").strip()
    if not nombre or not valor:
        raise ErrorUsuario("Ese campo está vacío.")
    origen, url = origen_de(con, "fila", fila_id)
    titulo = f"{fila.titulo.splitlines()[0][:120]} — {nombre}"
    return _insertar(rcon, rid, "texto", titulo=titulo, contenido=valor, origen=origen, origen_url=url)


def agregar_nota(con: sqlite3.Connection, rcon: sqlite3.Connection, nota_id: int, rid: int) -> int:
    nota = notas_srv.obtener(con, nota_id)
    if not nota.contenido.strip():
        raise ErrorUsuario("Esta nota está vacía.")
    origen, url = origen_de(con, "nota", nota_id)
    return _insertar(rcon, rid, "texto", titulo=nota.titulo, contenido=nota.contenido, formato=nota.formato,
                     origen=origen, origen_url=url)


def agregar_texto(con: sqlite3.Connection, rcon: sqlite3.Connection, texto: str, tipo: str, ref_id: int, titulo: str,
                  rid: int, pagina: int | None = None) -> int:
    """Un trozo de texto subrayado en la página (de una ficha, nota, imagen o PDF)."""
    texto = re.sub(r"[ \t ]+", " ", (texto or "").replace("\r", "")).strip()
    texto = re.sub(r"\n{3,}", "\n\n", texto)
    if not texto:
        raise ErrorUsuario("Primero subraya el texto que quieres guardar.")
    if len(texto) > 50000:
        raise ErrorUsuario("El texto subrayado es demasiado largo. Agrega la nota o la ficha completa.")
    origen, url = origen_de(con, tipo, ref_id, pagina)
    return _insertar(rcon, rid, "texto", titulo=" ".join((titulo or "").split())[:200], contenido=texto,
                     origen=origen, origen_url=url)


# ---------------------------------------------------------- imágenes y recortes
def leer_recorte(recorte: str) -> tuple[float, float, float, float] | None:
    """«x,y,ancho,alto» en fracciones de la imagen (0 a 1). Vacío = imagen completa."""
    if not (recorte or "").strip():
        return None
    try:
        x, y, w, h = (float(v) for v in recorte.split(","))
    except ValueError:
        raise ErrorUsuario("El recorte no es válido.") from None
    x, y = min(max(x, 0.0), 1.0), min(max(y, 0.0), 1.0)
    w, h = min(max(w, 0.0), 1.0 - x), min(max(h, 0.0), 1.0 - y)
    if w < 0.005 or h < 0.005:
        raise ErrorUsuario("El recorte es demasiado pequeño. Marca un rectángulo más grande.")
    return x, y, w, h


def imagen_fuente(con: sqlite3.Connection, tipo: str, ref_id: int, pagina: int | None = None,
                  recorte: tuple[float, float, float, float] | None = None, dpi: int = DPI_PDF) -> Image.Image:
    """La imagen original (o la página del PDF renderizada), recortada si se pide."""
    if tipo == "pdf":
        import pymupdf
        pdf = pdf_srv.obtener(con, ref_id)
        if not pagina or not 1 <= pagina <= pdf.paginas:
            raise ErrorUsuario(f"El PDF tiene {pdf.paginas} páginas: elige una entre 1 y {pdf.paginas}.")
        with pymupdf.open(pdf_srv.ruta(con, pdf)) as doc:
            pag = doc[pagina - 1]
            r = pag.rect
            clip = r if recorte is None else pymupdf.Rect(r.x0 + recorte[0] * r.width, r.y0 + recorte[1] * r.height,
                                                          r.x0 + (recorte[0] + recorte[2]) * r.width,
                                                          r.y0 + (recorte[1] + recorte[3]) * r.height)
            pix = pag.get_pixmap(dpi=dpi, clip=clip, alpha=False)
            return Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
    if tipo == "imagen":
        ruta = img_srv.ruta(con, img_srv.obtener(con, ref_id))
    elif tipo == "imagen_fila":
        ruta = imf.ruta(con, imf.obtener(con, ref_id))
    else:
        raise ErrorUsuario("Este elemento no es una imagen.")
    if not ruta.exists():
        raise NoEncontrado("El archivo de la imagen")
    with Image.open(ruta) as original:
        imagen = original.copy()
    if imagen.mode not in ("RGB", "RGBA", "L"):
        imagen = imagen.convert("RGBA" if "transparency" in imagen.info or imagen.mode in ("P", "LA") else "RGB")
    if recorte is not None:
        ancho, alto = imagen.size
        caja = (round(recorte[0] * ancho), round(recorte[1] * alto),
                round((recorte[0] + recorte[2]) * ancho), round((recorte[1] + recorte[3]) * alto))
        imagen = imagen.crop(caja)
    return imagen


def _titulo_imagen(con: sqlite3.Connection, tipo: str, ref_id: int, pagina: int | None) -> str:
    if tipo == "imagen":
        return img_srv.obtener(con, ref_id).titulo
    if tipo == "imagen_fila":
        return imf.obtener(con, ref_id).titulo
    return f"{pdf_srv.obtener(con, ref_id).titulo} — página {pagina}"


def agregar_imagen(con: sqlite3.Connection, rcon: sqlite3.Connection, tipo: str, ref_id: int, rid: int,
                   pagina: int | None = None, recorte: str = "", titulo: str = "") -> int:
    caja = leer_recorte(recorte)
    imagen = imagen_fuente(con, tipo, ref_id, pagina, caja)
    imagen.thumbnail((LADO_MAXIMO, LADO_MAXIMO))
    if imagen.mode == "RGBA":   # el PDF y Word no necesitan transparencia: fondo blanco
        fondo = Image.new("RGB", imagen.size, "white")
        fondo.paste(imagen, mask=imagen.getchannel("A"))
        imagen = fondo
    nombre = f"{uuid.uuid4().hex}.png"
    imagen.save(carpeta_imagenes() / nombre, "PNG", optimize=True)
    origen, url = origen_de(con, tipo, ref_id, pagina)
    titulo = " ".join((titulo or "").split()) or _titulo_imagen(con, tipo, ref_id, pagina) + (" (recorte)" if caja else "")
    return _insertar(rcon, rid, "imagen", titulo=titulo[:200], archivo=nombre, ancho=imagen.width, alto=imagen.height,
                     tamano="mediana" if caja else "grande", origen=origen, origen_url=url)


def ruta_imagen(nombre: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}\.png", nombre or ""):
        raise NoEncontrado("La imagen")
    return carpeta_imagenes() / nombre


def _borrar_imagen(nombre: str) -> None:
    try:
        ruta_imagen(nombre).unlink(missing_ok=True)
    except (ErrorUsuario, OSError):
        pass


# ------------------------------------------------------------------ exportar
def nombre_archivo(r: Recopilacion, extension: str) -> str:
    base = " ".join(re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", r.titulo).split()).strip(" .") or "Recopilacion"
    return f"{base[:80]}.{extension}"


CSS_PDF = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.4; color: #1f1b17; }
h1 { font-size: 20pt; margin: 0 0 4pt 0; color: #8f1d2c; }
h2.seccion { font-size: 15pt; color: #8f1d2c; margin: 16pt 0 4pt 0; padding-bottom: 2pt; border-bottom: 1px solid #d2ccc3; }
h3 { font-size: 12pt; margin: 12pt 0 3pt 0; }
h4, h5, h6 { font-size: 11pt; margin: 8pt 0 2pt 0; }
p { margin: 0 0 5pt 0; }
.portada { color: #6b645c; font-size: 9.5pt; }
.origen { color: #6b645c; font-size: 8.5pt; margin-top: 2pt; }
.dato { margin: 0 0 4pt 0; }
.comentario { background-color: #f7efda; padding: 5pt 7pt; margin: 4pt 0 6pt 0; }
.pie { color: #6b645c; font-size: 9pt; }
table { border-collapse: collapse; margin: 4pt 0 6pt 0; }
td, th { border: 1px solid #b9b2a8; padding: 2pt 4pt; font-size: 9.5pt; vertical-align: top; }
th { background-color: #f3f1ee; }
pre, code { font-family: monospace; font-size: 9pt; }
blockquote { margin: 4pt 0 4pt 12pt; color: #4a443e; }
"""


def documento_html(r: Recopilacion, src_imagen, logo: str | None = None) -> str:
    """El documento en HTML: lo usan la vista previa (navegador) y el PDF (PyMuPDF), así se ven igual."""
    e = html.escape
    partes = []
    if logo:
        partes.append(f'<p><img src="{e(logo)}" width="44" height="44"></p>')
    partes.append(f"<h1>{e(r.titulo)}</h1>")
    if r.descripcion:
        partes.append(f'<p>{e(r.descripcion)}</p>')
    fecha = datetime.now().strftime("%d/%m/%Y")
    partes.append(f'<p class="portada">{fecha} · {len([f for f in r.fragmentos if f.tipo != "seccion"])} fragmentos</p>')
    for fr in r.fragmentos:
        if fr.tipo == "seccion":
            partes.append(f'<h2 class="seccion">{e(fr.titulo)}</h2>')
            if fr.html_comentario:
                partes.append(f'<div class="comentario">{fr.html_comentario}</div>')
            continue
        if fr.titulo:
            partes.append(f"<h3>{e(fr.titulo)}</h3>")
        if fr.tipo == "campos":
            for nombre, valor in fr.campos:
                partes.append(f'<p class="dato"><b>{e(nombre)}:</b> {e(valor).replace(chr(10), "<br>")}</p>')
        elif fr.tipo == "texto":
            if fr.formato == "markdown":
                partes.append(fr.html)
            else:
                partes.append("".join(f"<p>{e(p).replace(chr(10), '<br>')}</p>" for p in re.split(r"\n\s*\n", fr.contenido)))
        elif fr.tipo == "imagen" and fr.archivo:
            ancho, alto = fr.medidas()
            partes.append(f'<p><img src="{e(src_imagen(fr.archivo))}" width="{ancho}" height="{alto}"></p>')
        if fr.comentario.strip():
            partes.append(f'<div class="comentario">{fr.html_comentario}</div>')
        if fr.origen:
            partes.append(f'<p class="origen">Fuente: {e(fr.origen)}</p>')
    return "\n".join(partes)


def exportar_pdf(r: Recopilacion, logo: Path | None = None) -> bytes:
    import pymupdf

    archivo = pymupdf.Archive(str(carpeta_imagenes()))
    if logo and logo.exists():
        archivo.add(pymupdf.Archive(str(logo.parent)))
    cuerpo = documento_html(r, lambda nombre: nombre, logo.name if logo and logo.exists() else None)
    historia = pymupdf.Story(html=cuerpo, user_css=CSS_PDF, archive=archivo)
    salida = io.BytesIO()
    escritor = pymupdf.DocumentWriter(salida)
    pagina = pymupdf.paper_rect("letter")
    zona = pagina + (54, 54, -54, -60)
    mas = True
    while mas:
        dispositivo = escritor.begin_page(pagina)
        mas, _ = historia.place(zona)
        historia.draw(dispositivo)
        escritor.end_page()
    escritor.close()
    # Numeración de páginas y título al pie
    with pymupdf.open("pdf", salida.getvalue()) as doc:
        total = doc.page_count
        for i, pag in enumerate(doc, start=1):
            pie = pymupdf.Rect(54, pagina.height - 42, pagina.width - 54, pagina.height - 24)
            pag.insert_textbox(pie, f"{r.titulo}  ·  página {i} de {total}", fontsize=8, color=(0.42, 0.39, 0.36),
                               align=pymupdf.TEXT_ALIGN_CENTER)
        doc.set_metadata({"title": r.titulo, "creator": "Biblioteca de Estudio"})
        return doc.tobytes(garbage=3, deflate=True)


# ---------------------------------------------------------------- Word (.docx)
def _md_tokens():
    return notas_srv._md  # mismo intérprete de Markdown que las notas (tablas incluidas)


def _runs(parrafo, hijos, negrita=False, cursiva=False) -> None:
    """Escribe los trozos de una línea de Markdown (negrita, cursiva, código, saltos)."""
    from docx.shared import Pt
    pila_b, pila_i = negrita, cursiva
    for t in hijos or []:
        if t.type == "strong_open":
            pila_b = True
        elif t.type == "strong_close":
            pila_b = negrita
        elif t.type == "em_open":
            pila_i = True
        elif t.type == "em_close":
            pila_i = cursiva
        elif t.type in ("softbreak", "hardbreak"):
            parrafo.add_run().add_break()
        elif t.type in ("text", "code_inline", "html_inline"):
            run = parrafo.add_run(t.content)
            run.bold, run.italic = pila_b or None, pila_i or None
            if t.type == "code_inline":
                run.font.name = "Courier New"
                run.font.size = Pt(9)


def _markdown_a_docx(doc, texto: str) -> None:
    tokens = _md_tokens().parse(texto)
    listas: list[str] = []
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if t.type == "heading_open":
            nivel = min(int(t.tag[1]) + 2, 9)
            p = doc.add_heading(level=nivel)
            _runs(p, tokens[i + 1].children)
            i += 3
            continue
        if t.type in ("bullet_list_open", "ordered_list_open"):
            listas.append("List Bullet" if t.type == "bullet_list_open" else "List Number")
        elif t.type in ("bullet_list_close", "ordered_list_close"):
            listas.pop()
        elif t.type == "paragraph_open":
            estilo = listas[-1] if listas else None
            if estilo and len(listas) > 1:
                estilo += f" {min(len(listas), 3)}"
            p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
            _runs(p, tokens[i + 1].children)
            i += 3
            continue
        elif t.type in ("fence", "code_block"):
            p = doc.add_paragraph()
            run = p.add_run(t.content.rstrip("\n"))
            run.font.name = "Courier New"
        elif t.type == "hr":
            doc.add_paragraph("―" * 20)
        elif t.type == "table_open":
            filas, j = [], i + 1
            while tokens[j].type != "table_close":
                if tokens[j].type == "tr_open":
                    filas.append([])
                elif tokens[j].type == "inline":
                    filas[-1].append((tokens[j].children, tokens[j - 1].type == "th_open"))
                j += 1
            ancho = max((len(f) for f in filas), default=0)
            if ancho:
                tabla = doc.add_table(rows=0, cols=ancho)
                tabla.style = "Table Grid"
                for f in filas:
                    if all(not h for h, _ in f) and all(e for _, e in f):
                        continue   # encabezado vacío (tablas pegadas sin encabezado)
                    celdas = tabla.add_row().cells
                    for k, (hijos, encabezado) in enumerate(f):
                        _runs(celdas[k].paragraphs[0], hijos, negrita=encabezado)
                doc.add_paragraph()
            i = j + 1
            continue
        i += 1


def _texto_a_docx(doc, texto: str) -> None:
    for bloque in re.split(r"\n\s*\n", texto.strip()):
        p = doc.add_paragraph()
        for k, linea in enumerate(bloque.split("\n")):
            if k:
                p.add_run().add_break()
            p.add_run(linea)


def exportar_docx(r: Recopilacion, logo: Path | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    vino, tenue = RGBColor(0x8F, 0x1D, 0x2C), RGBColor(0x6B, 0x64, 0x5C)
    doc = Document()
    seccion = doc.sections[0]
    seccion.page_width, seccion.page_height = Inches(8.5), Inches(11)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(seccion, lado, Inches(0.75))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    doc.core_properties.title = r.titulo

    if logo and logo.exists():
        doc.add_picture(str(logo), width=Pt(44))
    titulo = doc.add_heading(r.titulo, level=0)
    for run in titulo.runs:
        run.font.color.rgb = vino
    if r.descripcion:
        doc.add_paragraph(r.descripcion)
    p = doc.add_paragraph()
    run = p.add_run(f"{datetime.now():%d/%m/%Y} · {len([f for f in r.fragmentos if f.tipo != 'seccion'])} fragmentos")
    run.font.color.rgb, run.font.size = tenue, Pt(9)

    def comentario(fr: Fragmento) -> None:
        if not fr.comentario.strip():
            return
        inicio = len(doc.paragraphs)
        _markdown_a_docx(doc, fr.comentario)
        for par in doc.paragraphs[inicio:]:
            par.paragraph_format.left_indent = Pt(14)
            for run in par.runs:
                run.font.color.rgb = RGBColor(0x5A, 0x45, 0x10)
                run.italic = True

    for fr in r.fragmentos:
        if fr.tipo == "seccion":
            h = doc.add_heading(fr.titulo, level=1)
            for run in h.runs:
                run.font.color.rgb = vino
            comentario(fr)
            continue
        if fr.titulo:
            doc.add_heading(fr.titulo, level=2)
        if fr.tipo == "campos":
            for nombre, valor in fr.campos:
                par = doc.add_paragraph()
                par.add_run(f"{nombre}: ").bold = True
                for k, linea in enumerate(valor.split("\n")):
                    if k:
                        par.add_run().add_break()
                    par.add_run(linea)
        elif fr.tipo == "texto":
            (_markdown_a_docx if fr.formato == "markdown" else _texto_a_docx)(doc, fr.contenido)
        elif fr.tipo == "imagen" and fr.archivo:
            ruta = ruta_imagen(fr.archivo)
            if ruta.exists():
                ancho, _ = fr.medidas()
                doc.add_picture(str(ruta), width=Pt(ancho))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
        comentario(fr)
        if fr.origen:
            par = doc.add_paragraph()
            run = par.add_run(f"Fuente: {fr.origen}")
            run.font.size, run.font.color.rgb = Pt(8.5), tenue

    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()
