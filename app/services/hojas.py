"""Lectura de Excel (.xlsx/.xlsm) y CSV a una cuadrícula de textos.

Reglas para datos difíciles:
- Celdas combinadas: todas las celdas del rango toman el valor de la esquina superior izquierda.
- Fechas: se muestran como en Excel según su formato (p. ej. "6-7" que Excel convirtió en fecha).
- Números: sin ".0" sobrante; porcentajes con "%".
- Fórmulas: se usa el último valor calculado que guardó Excel.
- Imágenes incrustadas: se ignoran.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..errores import ErrorUsuario
from .almacenamiento import sin_tildes

EXTENSIONES_EXCEL = {".xlsx", ".xlsm"}
EXTENSIONES_CSV = {".csv", ".tsv", ".txt"}
EXTENSIONES = EXTENSIONES_EXCEL | EXTENSIONES_CSV
MAX_FILAS = 200_000


@dataclass
class Hoja:
    nombre: str
    filas: list[list[str]] = field(default_factory=list)

    @property
    def num_filas(self) -> int:
        return len(self.filas)

    @property
    def num_columnas(self) -> int:
        return max((len(f) for f in self.filas), default=0)


@dataclass
class Columna:
    indice: int          # 0-based dentro de la hoja
    letra: str           # "A", "B", …
    nombre: str          # encabezado (o "Columna C" si está vacío)
    ejemplos: list[str]  # primeros valores no vacíos
    llenas: int          # celdas con datos


# ----------------------------------------------------------------- celdas
_FECHA = re.compile(r"yyyy|yy|mmmm|mmm|mm|m|dd|d|hh|h|ss|s", re.IGNORECASE)


def _formatear_fecha(valor: date | datetime | time, formato: str) -> str:
    fmt = re.sub(r'\\(.)|"([^"]*)"', lambda m: m.group(1) or m.group(2) or "", formato or "")
    fmt = re.sub(r"\[[^\]]*\]", "", fmt.split(";")[0]).strip()
    if isinstance(valor, time):
        return valor.strftime("%H:%M" if valor.second == 0 else "%H:%M:%S")
    if not _FECHA.search(fmt) or fmt.lower() == "general":
        if isinstance(valor, datetime) and (valor.hour or valor.minute):
            return valor.strftime("%d/%m/%Y %H:%M")
        return valor.strftime("%d/%m/%Y")
    tiene_hora = bool(re.search(r"h", fmt, re.I))
    partes = {"yyyy": f"{valor.year}", "yy": f"{valor.year % 100:02d}",
              "dd": f"{valor.day:02d}", "d": str(valor.day),
              "mmmm": str(valor.month), "mmm": str(valor.month)}

    def token(m: re.Match) -> str:
        t = m.group(0).lower()
        if t in ("mm", "m"):
            # "m" junto a horas son minutos; si no, mes
            if tiene_hora and isinstance(valor, datetime) and re.search(r"h[^dmy]*$", fmt[: m.start()], re.I):
                return f"{valor.minute:02d}" if t == "mm" else str(valor.minute)
            return f"{valor.month:02d}" if t == "mm" else str(valor.month)
        if t in ("hh", "h") and isinstance(valor, datetime):
            return f"{valor.hour:02d}" if t == "hh" else str(valor.hour)
        if t in ("ss", "s") and isinstance(valor, datetime):
            return f"{valor.second:02d}"
        return partes.get(t, m.group(0))

    return _FECHA.sub(token, fmt)


def texto_celda(valor: object, formato: str = "General") -> str:
    """Valor de una celda tal como se ve en Excel."""
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, (datetime, date, time)):
        return _formatear_fecha(valor, formato)
    if isinstance(valor, (int, float)):
        if "%" in (formato or "") and isinstance(valor, (int, float)):
            return f"{valor * 100:.10g}%"
        if isinstance(valor, float):
            return str(int(valor)) if valor.is_integer() else f"{valor:.10g}"
        return str(valor)
    return str(valor)


# ----------------------------------------------------------------- lectura
def leer_archivo(ruta: Path, nombre_original: str = "") -> list[Hoja]:
    ext = Path(nombre_original or ruta.name).suffix.lower()
    if ext == ".xls":
        raise ErrorUsuario("Los archivos .xls (Excel antiguo) no son compatibles. Ábrelo en Excel y guárdalo como .xlsx.")
    if ext in EXTENSIONES_EXCEL:
        return _leer_excel(ruta)
    if ext in EXTENSIONES_CSV:
        return [_leer_csv(ruta, Path(nombre_original or ruta.name).stem)]
    raise ErrorUsuario("Formato no compatible. Sube un archivo de Excel (.xlsx) o CSV.")


def _leer_excel(ruta: Path) -> list[Hoja]:
    try:
        libro = load_workbook(ruta, data_only=True)
    except Exception as e:  # archivo dañado, protegido con contraseña, etc.
        raise ErrorUsuario("No se pudo abrir el Excel. Comprueba que no esté dañado ni protegido con contraseña.") from e
    hojas = []
    try:
        for ws in libro.worksheets:
            if ws.sheet_state != "visible":
                continue
            max_fila, max_col = min(ws.max_row, MAX_FILAS), ws.max_column
            cuadricula = [[""] * max_col for _ in range(max_fila)]
            for fila in ws.iter_rows(min_row=1, max_row=max_fila, max_col=max_col):
                for celda in fila:
                    if celda.value is not None:
                        cuadricula[celda.row - 1][celda.column - 1] = texto_celda(celda.value, celda.number_format)
            for rango in ws.merged_cells.ranges:  # celdas combinadas: se repite el valor
                valor = cuadricula[rango.min_row - 1][rango.min_col - 1] if rango.min_row <= max_fila else ""
                for r in range(rango.min_row, min(rango.max_row, max_fila) + 1):
                    for c in range(rango.min_col, min(rango.max_col, max_col) + 1):
                        cuadricula[r - 1][c - 1] = valor
            hojas.append(Hoja(ws.title, _recortar(cuadricula)))
    finally:
        libro.close()
    if not any(h.filas for h in hojas):
        raise ErrorUsuario("El archivo no tiene datos.")
    return hojas


def _leer_csv(ruta: Path, nombre: str) -> Hoja:
    crudo = ruta.read_bytes()
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = crudo.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    muestra = texto[:20000]
    try:
        dialecto: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
    except csv.Error:
        separador = ";" if muestra.count(";") > muestra.count(",") else ","

        class dialecto(csv.excel):  # type: ignore[no-redef]
            delimiter = separador
    filas = [list(f) for f in csv.reader(io.StringIO(texto), dialecto)][:MAX_FILAS]
    if not filas:
        raise ErrorUsuario("El archivo CSV está vacío.")
    ancho = max(len(f) for f in filas)
    return Hoja(nombre, _recortar([f + [""] * (ancho - len(f)) for f in filas]))


def _recortar(cuadricula: list[list[str]]) -> list[list[str]]:
    """Quita filas y columnas vacías del final."""
    while cuadricula and not any(v.strip() for v in cuadricula[-1]):
        cuadricula.pop()
    if not cuadricula:
        return []
    ultima = max((i for f in cuadricula for i, v in enumerate(f) if v.strip()), default=-1)
    return [f[: ultima + 1] for f in cuadricula]


# --------------------------------------------------------------- encabezados
def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", sin_tildes(t)).strip().lower()


def detectar_encabezado(hoja: Hoja) -> int:
    """Fila (1-based) que parece el encabezado: la primera bastante llena entre las primeras 30."""
    primeras = hoja.filas[:30]
    # Se cuentan valores distintos: un título combinado repite el mismo texto en varias celdas
    llenas = [len({v.strip() for v in f if v.strip()}) for f in primeras]
    if not llenas or max(llenas) == 0:
        return 1
    umbral = max(2, int(max(llenas) * 0.5)) if max(llenas) > 1 else 1
    for i, n in enumerate(llenas):
        if n >= umbral:
            return i + 1
    return 1


def parece_subencabezado(hoja: Hoja, fila_encabezado: int) -> bool:
    """¿La fila siguiente repite (al menos la mitad de) los encabezados? Pasa en algunos Excel."""
    if fila_encabezado >= hoja.num_filas:
        return False
    arriba, abajo = hoja.filas[fila_encabezado - 1], hoja.filas[fila_encabezado]
    pares = [(a, b) for a, b in zip(arriba, abajo) if b.strip()]
    return bool(pares) and sum(_norm(a) == _norm(b) for a, b in pares) / len(pares) >= 0.5


def columnas(hoja: Hoja, fila_encabezado: int, dos_filas: bool = False) -> list[Columna]:
    """Columnas disponibles (las que tienen encabezado o algún dato), con nombres únicos."""
    inicio = fila_encabezado + (1 if dos_filas else 0)
    datos = hoja.filas[inicio:]
    arriba = hoja.filas[fila_encabezado - 1] if hoja.filas else []
    abajo = hoja.filas[fila_encabezado] if dos_filas and fila_encabezado < hoja.num_filas else []
    resultado: list[Columna] = []
    usados: dict[str, int] = {}
    for i in range(hoja.num_columnas):
        a = " ".join((arriba[i] if i < len(arriba) else "").split())
        b = " ".join((abajo[i] if i < len(abajo) else "").split())
        valores = [f[i] for f in datos if i < len(f) and f[i].strip()]
        if not (a or b or valores):
            continue
        nombre = a or b or f"Columna {get_column_letter(i + 1)}"
        if a and b and _norm(a) != _norm(b):
            nombre = f"{a} · {b}"
        clave = _norm(nombre)
        if clave in usados:  # encabezados repetidos (p. ej. celdas combinadas)
            usados[clave] += 1
            nombre = f"{nombre} ({usados[clave]})"
        else:
            usados[clave] = 1
        resultado.append(Columna(i, get_column_letter(i + 1), nombre[:120], [v[:80] for v in valores[:3]], len(valores)))
    return resultado


def columna_titulo_sugerida(cols: list[Columna]) -> int | None:
    preferidas = ("nombre", "titulo", "medicamento", "termino", "concepto", "tema", "estructura", "farmaco")
    for c in cols:
        if any(p in _norm(c.nombre) for p in preferidas):
            return c.indice
    return cols[0].indice if cols else None


def registros(hoja: Hoja, fila_encabezado: int, dos_filas: bool, indices: list[int]) -> list[tuple[int, list[str]]]:
    """Filas de datos (número de fila en el archivo, valores de las columnas elegidas). Omite filas vacías."""
    inicio = fila_encabezado + (1 if dos_filas else 0)
    salida = []
    for n, fila in enumerate(hoja.filas[inicio:], start=inicio + 1):
        valores = [fila[i] if i < len(fila) else "" for i in indices]
        if any(v.strip() for v in valores):
            salida.append((n, valores))
    return salida


def ancho_columna(nombre: str, valores: list[str]) -> int:
    """Ancho sugerido (px) según el contenido: columnas cortas angostas, textos largos más anchos."""
    largos = sorted(max((len(linea) for linea in v.splitlines()), default=0) for v in valores if v.strip())
    p90 = largos[max(int(len(largos) * 0.9) - 1, 0)] if largos else 0
    palabra = max((len(w) for v in valores[:500] for w in v.split()), default=0)
    return int(max(110, min(360, max(p90 * 7.0 + 40, min(len(nombre), 24) * 7.6 + 40, min(palabra * 7.4 + 36, 360)))))
