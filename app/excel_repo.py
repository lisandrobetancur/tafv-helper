"""Lectura (solo lectura) del Excel que funciona como base de datos.

Nunca se escribe sobre el archivo: se abre, se copia lo necesario a memoria
y se cierra.

Los campos NO están fijos en el código: se descubren a partir de los
encabezados de la hoja, dentro del rango de columnas configurado
(`columns`, por defecto "B:O"). Si el archivo trae menos o más columnas
dentro de ese rango, la aplicación se adapta sola.
"""
from __future__ import annotations

import re
import threading
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

SHEET_ESTUDIO = "Estudio homeopatia"
SHEET_FORMULACION = "Formulación"
DEFAULT_COLUMNS = "B:O"
HEADER_ROW = 1
RESERVED = {"id", "row", "titulo", "adjuntos", "_norm"}


# --------------------------------------------------------------- utilidades
def normalize(text: str) -> str:
    """Minúsculas, sin acentos y con espacios colapsados (para buscar/filtrar)."""
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", text).strip().lower()


def slugify(text: str, sep: str = "-") -> str:
    return re.sub(r"[^a-z0-9]+", sep, normalize(text)).strip(sep) or "campo"


_DATE_TOKEN = re.compile(r"yyyy|yy|mmmm|mmm|mm|m|dd|d", re.IGNORECASE)


def _render_date(value: date, number_format: str) -> str:
    """Reconstruye el texto digitado cuando Excel convirtió un valor como "6-7" en fecha.

    En este archivo no hay fechas reales: todo es texto o número, así que se
    devuelve la fecha tal como la muestra su formato (d-m -> "6-7", m/d/yy -> "6/7/25").
    """
    fmt = re.sub(r'\\(.)|"([^"]*)"', lambda m: m.group(1) or m.group(2), number_format or "")
    fmt = re.sub(r"\[[^\]]*\]", "", fmt.split(";")[0])  # [$-409], condiciones, etc.
    if not _DATE_TOKEN.search(fmt) or fmt.lower() == "general":
        fmt = "d-m"  # lo más habitual al digitar "6-7"

    def token(m: re.Match) -> str:
        return {
            "d": str(value.day), "dd": f"{value.day:02d}",
            "m": str(value.month), "mm": f"{value.month:02d}",
            "mmm": str(value.month), "mmmm": str(value.month),
            "yy": f"{value.year % 100:02d}", "yyyy": str(value.year),
        }[m.group(0).lower()]

    return _DATE_TOKEN.sub(token, fmt)


def cell_text(cell) -> str:
    """Texto de una celda tal como se ve en Excel (sin alterar el contenido)."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return _render_date(value, cell.number_format)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_columns(spec: str) -> tuple[int, int | None]:
    """"B:O" -> (2, 15); "B:" -> (2, None) = hasta la última columna con datos."""
    first, _, last = (spec or DEFAULT_COLUMNS).replace(" ", "").upper().partition(":")
    lo = column_index_from_string(first)
    hi = column_index_from_string(last) if last else None
    if hi is not None and hi < lo:
        lo, hi = hi, lo
    return lo, hi


def clean_label(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _column_width(label: str, values: list[str]) -> int:
    """Ancho (px) según el contenido: columnas cortas angostas, textos largos más anchos."""
    lengths = sorted(max((len(line) for line in v.splitlines()), default=0) for v in values if v.strip())
    p90 = lengths[max(int(len(lengths) * 0.9) - 1, 0)] if lengths else 0
    longest_word = max((len(w) for v in values for w in v.split()), default=0)
    content = p90 * 7.0 + 40
    header = min(len(label), 22) * 7.4 + 36
    floor = min(longest_word * 7.4 + 36, 360)  # que no se parta una palabra
    return int(max(96, min(360, max(content, header, floor))))


# ------------------------------------------------------------------ lectura
class ExcelRepo:
    def __init__(self, path: Path, columns: str = DEFAULT_COLUMNS):
        self.path = Path(path)
        self.columns = columns
        self._lock = threading.Lock()
        self.fields: list[dict] = []      # {key, label, col, letter, width}
        self.filters: list[dict] = []     # {key, label, placeholder, sources}
        self.records: list[dict] = []
        self.by_id: dict[str, dict] = {}
        self.formulacion: dict = {"rows": []}
        self.loaded_at: datetime | None = None
        self.load()

    def load(self) -> None:
        if not self.path.exists():
            raise FileNotFoundError(f"No se encontró el Excel en: {self.path}")
        wb = openpyxl.load_workbook(self.path, data_only=True)
        try:
            fields, records = self._read_estudio(wb[SHEET_ESTUDIO], self.columns)
            formulacion = self._read_formulacion(wb[SHEET_FORMULACION])
        finally:
            wb.close()
        filters = self._build_filters(fields)
        hidden = {k for flt in filters if flt.get("hide_column") for k in flt["sources"]}
        for f in fields:  # su filtro propio las sustituye en la tabla de resultados
            f["hide_in_filters"] = f["key"] in hidden
        with self._lock:
            self.fields, self.records, self.filters = fields, records, filters
            self.by_id = {r["id"]: r for r in records}
            self.formulacion = formulacion
            self.loaded_at = datetime.now()

    # -- hoja principal ------------------------------------------------------
    @staticmethod
    def _read_estudio(ws, columns: str) -> tuple[list[dict], list[dict]]:
        lo, hi = parse_columns(columns)
        hi = min(hi or ws.max_column, ws.max_column)
        cols = range(lo, hi + 1)

        def row_texts(r: int) -> list[str]:
            return [cell_text(ws.cell(r, c)) for c in cols]

        head1, head2 = row_texts(HEADER_ROW), row_texts(HEADER_ROW + 1)

        # ¿La fila 2 es un subencabezado? (repite los encabezados de la fila 1)
        filled2 = [(a, b) for a, b in zip(head1, head2) if b.strip()]
        sub = bool(filled2) and sum(normalize(a) == normalize(b) for a, b in filled2) / len(filled2) >= 0.5
        first_data = HEADER_ROW + (2 if sub else 1)

        rows = [(r, row_texts(r)) for r in range(first_data, ws.max_row + 1)]
        rows = [(r, t) for r, t in rows if any(x.strip() for x in t)]

        # Una columna existe si tiene encabezado o algún dato
        fields: list[dict] = []
        used_keys: set[str] = set()
        for i, c in enumerate(cols):
            has_data = any(t[i].strip() for _, t in rows)
            top = clean_label(head1[i])
            bottom = clean_label(head2[i]) if sub else ""
            if not (top or bottom or has_data):
                continue
            label = top or bottom or f"Columna {get_column_letter(c)}"
            if top and bottom and normalize(top) != normalize(bottom):
                label = f"{top} · {bottom}"
            key = slugify(label, "_")
            if key in RESERVED:
                key += "_campo"
            base, n = key, 2
            while key in used_keys:
                key, n = f"{base}_{n}", n + 1
            used_keys.add(key)
            fields.append({"key": key, "label": label, "col": c, "idx": i,
                           "letter": get_column_letter(c)})

        for f in fields:
            f["width"] = _column_width(f["label"], [t[f["idx"]] for _, t in rows])
        title_key = ExcelRepo._title_key(fields)

        records: list[dict] = []
        used_ids: set[str] = set()
        for r, texts in rows:
            values = {f["key"]: texts[f["idx"]] for f in fields}
            title = values.get(title_key, "").strip() if title_key else ""
            if not title:  # sin nombre principal: primer valor no vacío
                title = next((v.strip() for v in values.values() if v.strip()), "")
            base = slugify(title) if title else f"fila-{r}"
            rid, n = base, 2
            while rid in used_ids:
                rid, n = f"{base}-{n}", n + 1
            used_ids.add(rid)
            records.append({
                "id": rid, "row": r,
                "titulo": title or f"(sin nombre · fila {r})",
                **values,
                "_norm": {k: normalize(v) for k, v in values.items()},
            })
        for f in fields:
            f.pop("idx")
        return fields, records

    @staticmethod
    def _title_key(fields: list[dict]) -> str | None:
        for pattern in ("medicamento", "latin", "nombre"):
            for f in fields:
                if pattern in normalize(f["label"]):
                    return f["key"]
        return fields[0]["key"] if fields else None

    @staticmethod
    def _build_filters(fields: list[dict]) -> list[dict]:
        """Filtros según los encabezados presentes; si no hay ninguno reconocible, búsqueda general."""
        def pick(*patterns: str) -> list[str]:
            return [f["key"] for f in fields if any(p in normalize(f["label"]) for p in patterns)]

        filters = []
        meds = pick("medicamento", "latin")
        if meds:
            filters.append({"key": "medicamento", "label": "Medicamento",
                            "placeholder": "Nombre del medicamento…", "sources": meds})
        otros = pick("otros sintoma")
        sint = [k for k in pick("sintoma") if k not in otros]
        if sint:
            filters.append({"key": "sintomas", "label": "Síntomas",
                            "placeholder": "Escribe un síntoma…", "sources": sint})
        if otros:
            filters.append({"key": "otros_sintomas", "label": "Otros síntomas",
                            "placeholder": "Escribe otro síntoma…", "sources": otros,
                            "hide_column": True})
        if not filters and fields:
            filters.append({"key": "buscar", "label": "Buscar",
                            "placeholder": "Buscar en todos los campos…",
                            "sources": [f["key"] for f in fields]})
        return filters

    # -- hoja Formulación ----------------------------------------------------
    @staticmethod
    def _read_formulacion(ws) -> dict:
        coords = [c for r in ws.iter_rows() for c in r if cell_text(c).strip()]
        if not coords:
            return {"rows": []}
        r0, r1 = min(c.row for c in coords), max(c.row for c in coords)
        c0, c1 = min(c.column for c in coords), max(c.column for c in coords)

        spans: dict[tuple[int, int], tuple[int, int]] = {}
        covered: set[tuple[int, int]] = set()
        for rng in ws.merged_cells.ranges:
            spans[(rng.min_row, rng.min_col)] = (
                rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
            for rr in range(rng.min_row, rng.max_row + 1):
                for cc in range(rng.min_col, rng.max_col + 1):
                    if (rr, cc) != (rng.min_row, rng.min_col):
                        covered.add((rr, cc))

        rows = []
        for r in range(r0, r1 + 1):
            cells = []
            for c in range(c0, c1 + 1):
                if (r, c) in covered:
                    continue
                cell = ws.cell(r, c)
                rowspan, colspan = spans.get((r, c), (1, 1))
                cells.append({"text": cell_text(cell), "rowspan": rowspan, "colspan": colspan,
                              "bold": bool(cell.font and cell.font.b)})
            if any(x["text"].strip() for x in cells):  # omite filas totalmente vacías
                rows.append(cells)
        return {"rows": rows}

    # ---------------------------------------------------------------- consulta
    def search(self, params: dict[str, str]) -> list[dict]:
        """Filtra sin modificar los registros ("contiene", sin acentos ni mayúsculas)."""
        wanted = []
        for flt in self.filters:
            nv = normalize(params.get(flt["key"], ""))
            if nv:
                wanted.append((flt["sources"], nv))
        return [
            rec for rec in self.records
            if all(any(nv in rec["_norm"][f] for f in sources) for sources, nv in wanted)
        ]

    def suggest(self, key: str, q: str, limit: int = 12) -> list[str]:
        """Sugerencias de autocompletado: líneas de texto de los registros que contienen `q`."""
        flt = next((f for f in self.filters if f["key"] == key), None)
        if flt is None:
            raise KeyError(key)
        nq = normalize(q)
        seen: dict[str, str] = {}
        for rec in self.records:
            for field in flt["sources"]:
                for line in rec[field].splitlines():
                    text = line.strip().lstrip("-•*· ").strip()
                    if text and nq in normalize(text):
                        seen.setdefault(normalize(text), text[:80])
        return sorted(seen.values(), key=lambda t: (not normalize(t).startswith(nq), normalize(t)))[:limit]
