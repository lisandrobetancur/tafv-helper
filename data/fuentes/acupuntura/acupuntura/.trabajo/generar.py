"""Genera notas/*.md del componente Acupuntura a partir de hojas/Acupuntura.xlsx (literal + tildes)."""
import re
import sys
from pathlib import Path

import openpyxl

BASE = Path(sys.argv[1])
ws = openpyxl.load_workbook(BASE / "hojas/Acupuntura.xlsx", data_only=True).active

# Solo tildes y correcciones declaradas (palabra completa, respeta mayúsculas)
TILDES = {
    "ORGANOS": "ÓRGANOS", "VISCERAS": "VÍSCERAS", "FRIO": "FRÍO", "frio": "frío", "Frio": "Frío",
    "frios": "fríos", "rio": "río", "TECNICA": "TÉCNICA", "tecnicas": "técnicas", "PATOGENOS": "PATÓGENOS",
    "Inflamacion": "Inflamación", "Vertigo": "Vértigo", "Sincope": "Síncope", "medulas": "médulas",
    "SEPARACION": "SEPARACIÓN", "REUNION": "REUNIÓN", "TELURICA": "TELÚRICA", "ACELERACION": "ACELERACIÓN",
    "CIRCULACION": "CIRCULACIÓN", "TORACICA": "TORÁCICA", "ENERGIAS": "ENERGÍAS", "Sí": "Si",
    # cambios de letras (declarados en plan.json → correcciones)
    "DEFICIENCIE": "DEFICIENCIA", "aguada": "aguda",
}


def fix(s: str) -> str:
    s = s.replace("è", "→")  # flecha Wingdings perdida (declarada en correcciones)
    for a, b in TILDES.items():
        s = re.sub(rf"(?<![\wÁÉÍÓÚáéíóúñÑ]){re.escape(a)}(?![\wÁÉÍÓÚáéíóúñÑ])", b, s)
    return s


def v(coord: str, sep: str = "\n") -> str:
    x = ws[coord].value
    if x is None:
        return ""
    lineas = [l.strip() for l in str(x).split("\n")]
    return fix(sep.join(l for l in lineas if l))


def fila_md(celdas):
    return "| " + " | ".join(celdas) + " |"


def tabla(filas, encabezado=None):
    n = len(filas[0])
    out = [fila_md(encabezado) if encabezado else "|" + " |" * n, "|" + "---|" * n]
    out += [fila_md(f) for f in filas]
    return "\n".join(out).replace("|  |", "| |").replace("|  |", "| |")


def rango(cols, r0, r1):
    return [[v(f"{c}{r}", " ") for c in cols] for r in range(r0, r1 + 1)]


def lineas(col, r0, r1):
    return "\n".join(v(f"{col}{r}") for r in range(r0, r1 + 1) if v(f"{col}{r}"))


notas = {}

# TÉCNICA ZANG-FU (B136:R159; la columna B es el encabezado, va en el título)
cols = list("CDEFGHIJKLMNOPQR")
notas["TÉCNICA ZANG-FU - Acupuntura"] = tabla(rango(cols, 137, 159), [v(f"{c}136", " ") for c in cols])

# LUO (C161:F183)
notas["LUO - Acupuntura"] = tabla(rango("CDEF", 163, 183), [v(f"{c}162") for c in "CDEF"])

# YUAN (G161:J168)
notas["YUAN - Acupuntura"] = v("G162") + "\n\n" + tabla([[v(f"G{r}"), v(f"H{r}")] for r in range(163, 169)])

# MU (P162:Q173)
notas["MU - Acupuntura"] = tabla(rango("PQ", 163, 173))

# TÉCNICA DE PLANOS (B186 + B188:O201)
cols = list("BCDEFGHIJKLMNO")
notas["TÉCNICA DE PLANOS - Acupuntura"] = tabla(rango(cols, 189, 201), [v(f"{c}188") for c in cols])

# HUI (B222:C235) y XI (B236:C251)
notas["HUI - Acupuntura"] = tabla(rango("BC", 223, 235))
notas["XI - Acupuntura"] = tabla(rango("BC", 237, 251))

# HE (D222:E249)
notas["HE - Acupuntura"] = tabla(rango("DE", 223, 249))

# Columna F: 8 técnicas BA FENG y recuadros de color debajo
f231 = [l.strip() for l in ws["F231"].value.split("\n")]
ba_feng = [
    v("F223"), "   " + v("F224"),
    v("F225"), v("F226"), v("F227"), v("F228"), v("F229"), v("F230"),
    fix(f231[0]),
] + ["   " + fix(l) for l in f231[1:]]
notas["8 técnicas BA FENG - Acupuntura"] = "\n".join(ba_feng)
notas["Dolor de vísceras - Acupuntura"] = lineas("F", 233, 239)
notas["Dolor de órganos - Acupuntura"] = lineas("F", 241, 245)
notas["Enfermedad interna - Acupuntura"] = lineas("F", 247, 255)
notas["Técnica de LUO - Acupuntura"] = lineas("F", 257, 261)

# PATÓGENOS EXTERNOS (J222:M278): una nota por recuadro (grupo de color; FRÍO/VIENTO/HUMEDAD/FLEMA son
# encabezados propios). Título = encabezado (o primer renglón) + « (Patógenos externos)».
grupos = [(223, 226), (227, 230), (231, 237), (238, 248), (249, 252), (253, 254), (255, 257), (258, 263),
          (264, 267), (268, 270), (271, 273), (274, 275), (276, 278)]
CASO = {"FRÍO": "Frío", "VIENTO": "Viento", "HUMEDAD": "Humedad", "FLEMA": "Flema"}
PATOGENOS = []
for r0, r1 in grupos:
    renglones = [v(f"J{r}") for r in range(r0, r1 + 1)]
    enc = renglones[0]
    enc = CASO.get(enc, enc).rstrip(":").strip()
    enc = enc[0].upper() + enc[1:]
    titulo = f"{enc} (Patógenos externos) - Acupuntura"
    notas[titulo] = "\n".join(renglones[1:])
    PATOGENOS.append((titulo, f"J{r0}:M{r1}"))
(BASE / "notas").mkdir(exist_ok=True)
for titulo, contenido in notas.items():
    (BASE / "notas" / f"{titulo}.md").write_text(contenido.strip() + "\n", encoding="utf-8")
    print("escrita:", titulo)
(BASE / ".trabajo/patogenos.json").write_text(__import__("json").dumps(PATOGENOS, ensure_ascii=False))
