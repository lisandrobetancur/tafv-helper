"""Copia corregida de la tabla: tildes, letras cambiadas, palabras pegadas, mayúscula inicial,
saltos que cortan frase. Escribe correcciones.json con los cambios de letras."""
import json, re, sys, unicodedata
from pathlib import Path
import openpyxl

BASE = Path(sys.argv[1])
ORIG = BASE / "hojas/Estudio homeopatia.xlsx"
DEST = BASE / "tablas/Estudio homeopatia.xlsx"

def st(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")

T = """sintomas:síntomas sintoma:síntoma latin:latín forunculo:forúnculo sensacion:sensación aversion:aversión
progresion:progresión rapido:rápido rapida:rápida rapidos:rápidos termica:térmica patologico:patológico
patologicos:patológicos patologica:patológica tirania:tiranía agravacion:agravación frio:frío fria:fría frios:fríos
frias:frías humedo:húmedo humedos:húmedos constitucion:constitución carbon:carbón vertigo:vértigo oseas:óseas
oseo:óseo oseos:óseos especificos:específicos emaciacion:emaciación accion:acción frustracion:frustración
glandulas:glándulas glandula:glándula vision:visión oido:oído oidos:oídos congestion:congestión cefalica:cefálica
pulsatiles:pulsátiles pulsatil:pulsátil excitacion:excitación exposicion:exposición pronostico:pronóstico
supresion:supresión musculoesqueletico:musculoesquelético musculoesqueletica:musculoesquelética cafe:café
histerico:histérico animo:ánimo mejoria:mejoría distraccion:distracción toxica:tóxica toxico:tóxico toxicas:tóxicas
endocrinologia:endocrinología adenopatias:adenopatías publico:público fetida:fétida fetidas:fétidas fetido:fétido
fistulas:fístulas transpiracion:transpiración escalofrios:escalofríos supuracion:supuración maduracion:maduración
expulsion:expulsión cicatrizacion:cicatrización obstinacion:obstinación cartilagos:cartílagos fragiles:frágiles
fragil:frágil inorganico:inorgánico oxidos:óxidos organico:orgánico cutanea:cutánea cutaneas:cutáneas cutaneo:cutáneo
cutaneos:cutáneos paralisis:parálisis toracico:torácico toracica:torácica facilmente:fácilmente encias:encías
sudoracion:sudoración flacida:flácida palida:pálida palido:pálido ulceracion:ulceración coagulos:coágulos
colicos:cólicos colico:cólico colera:cólera asimetrica:asimétrica neuralgicos:neurálgicos denticion:dentición
nauseas:náuseas vomito:vómito vomitos:vómitos simultaneos:simultáneos ulceras:úlceras ulcera:úlcera
intoxicacion:intoxicación estres:estrés compañia:compañía dificil:difícil dificiles:difíciles periodicas:periódicas
anemico:anémico alergica:alérgica consolacion:consolación neurologico:neurológico espasmodico:espasmódico
espasmodica:espasmódica fisicamente:físicamente fisico:físico fisica:física impresion:impresión neuroticos:neuróticos
constriccion:constricción opresion:opresión emocion:emoción torax:tórax organos:órganos excoriacion:excoriación
espontaneos:espontáneos espontaneo:espontáneo facil:fácil energia:energía clonicos:clónicos
repertorizacion:repertorización pigmentacion:pigmentación vacio:vacío estomago:estómago acidos:ácidos acido:ácido
irritacion:irritación inflamacion:inflamación menstruacion:menstruación indigestion:indigestión atencion:atención
irresolucion:irresolución docil:dócil tardia:tardía melancolia:melancolía melancolico:melancólico depresion:depresión
melodias:melodías homeopatia:homeopatía colericos:coléricos hepaticos:hepáticos hepatico:hepático
caracteristico:característico caracteristicos:característicos caracteristica:característica caracter:carácter
tipicas:típicas tipica:típica tipico:típico descomposicion:descomposición posicion:posición liquida:líquida
liquido:líquido liquidos:líquidos cadaverico:cadavérico despues:después dermatologia:dermatología atopicas:atópicas
atopica:atópica exacerbacion:exacerbación epoca:época reaccion:reacción aspera:áspera preparacion:preparación
septica:séptica septicos:sépticos debil:débil debiles:débiles ortografia:ortografía cilindrico:cilíndrico
psicoticos:psicóticos produccion:producción crepusculo:crepúsculo cronico:crónico cronica:crónica cronicas:crónicas
perdida:pérdida confusion:confusión linfatico:linfático repeticion:repetición ningun:ningún estan:están
sindromes:síndromes sindrome:síndrome subito:súbito propension:propensión musica:música oncologicos:oncológicos
habitos:hábitos dilucion:dilución critica:crítica liberacion:liberación polipos:pólipos apatia:apatía
postracion:postración vacunacion:vacunación infeccion:infección asintomatico:asintomático vias:vías
recuperacion:recuperación silice:sílice vesicula:vesícula vesiculas:vesículas apatico:apático timido:tímido
perimenopausicos:perimenopáusicos secrecion:secreción melicerico:melicérico indecision:indecisión
constipacion:constipación anoxicas:anóxicas inhalacion:inhalación ingestion:ingestión lipidica:lipídica
picazon:picazón erupcion:erupción mecanicos:mecánicos gastricos:gástricos gastrica:gástrica medico:médico
adinamicos:adinámicos adinamico:adinámico contusion:contusión negacion:negación hiperico:hipérico
herpeticos:herpéticos postquirurgicos:postquirúrgicos caidas:caídas tetanos:tétanos sintomatica:sintomática
acumulacion:acumulación cardiopatias:cardiopatías rigidas:rígidas laringea:laríngea traqueitis:traqueítis
comio:comió cayo:cayó digestion:digestión corazon:corazón util:útil eligio:eligió recibio:recibió siguio:siguió
conclusion:conclusión articulo:artículo podria:podría criticos:críticos dia:día dias:días mejoro:mejoró
recupero:recuperó permanecia:permanecía valoracion:valoración funcion:función translucido:translúcido
reumaticos:reumáticos reumaticas:reumáticas esqueletico:esquelético neumonia:neumonía pleuresia:pleuresía
tumbandose:tumbándose ademas:además minimo:mínimo sosten:sostén presion:presión glucosidos:glucósidos
citotoxicos:citotóxicos heterosidos:heterósidos region:región violaceo:violáceo violacea:violácea
paradojica:paradójica evolucion:evolución hinchazon:hinchazón prevencion:prevención postraumaticas:postraumáticas
postraumatico:postraumático exaltacion:exaltación imaginacion:imaginación separacion:separación
agudizacion:agudización timpano:tímpano aparicion:aparición miccion:micción friccion:fricción aplicacion:aplicación
decubito:decúbito salivacion:salivación disentericas:disentéricas destruccion:destrucción sintomatologia:sintomatología
climatico:climático musculos:músculos musculo:músculo quemazon:quemazón irradiacion:irradiación depuracion:depuración
tonico:tónico estimulacion:estimulación acetico:acético mucopolisacaridos:mucopolisacáridos penfigos:pénfigos
bronconeumopatias:bronconeumopatías tambien:también disfonia:disfonía tapon:tapón impetigo:impétigo
neuropatia:neuropatía piromano:pirómano nefritico:nefrítico antimitotico:antimitótico toxicologico:toxicológico
toxicologia:toxicología manifestacion:manifestación rosaceas:rosáceas atras:atrás agitacion:agitación
resolucion:resolución mejoria:mejoría perfeccion:perfección congestion:congestión basica:básica distension:distensión
hepatica:hepática herpeticas:herpéticas sofocacion:sofocación inutil:inútil triangulo:triángulo"""
L = """ambibalente:ambivalente Sudoracio:Sudoración condicmentados:condimentados fluorathum:fluoratum inico:inicio
estricinina:estricnina delgadoz:delgados Chamomilia:Chamomilla entimecimiento:entumecimiento colpaso:colapso
postracción:postración intesa:intensa decepsiones:decepciones senstimiento:sentimiento afecicones:afecciones
mejillaz:mejillas avercion:aversión Ranuncolosido:Ranunculósido trastosnos:trastornos graos:grasos freco:fresco
mentruacio:menstruación dispesia:dispepsia aliemntos:alimentos calusosa:calurosa intolerania:intolerancia
retrasi:retraso espaciial:especial induracin:induración extrama:extrema doloros:dolorosa viajo:viejo clvada:clavada
crucificcion:crucifixión endotoxiina:endotoxina gonoccocica:gonocócica refrewscantes:refrescantes
amarilentas:amarillentas posicio:posición doloroidos:doloridos octurna:nocturna inmuopatología:inmunopatología
lilbros:libros calquiera:cualquiera emacion:emaciación Temperamiento:Temperamento tubeculínico:tuberculínico
suceptible:susceptible blafesma:blasfema agoramiento:agotamiento Periocidad:Periodicidad silece:sílice
contraacturas:contracturas caracteriticas:características lilamento:filamento cuteneas:cutáneas placton:plancton
NERVISO:NERVIOSO Nervisos:Nerviosos cinestosis:cinetosis cabellido:cabelludo neuseas:náuseas viajer:viajar
etleta:atleta alimenton:alimentos eruptos:eructos Eruptos:Eructos sesquiterrpenos:sesquiterpenos prevencio:prevención
creneales:craneales Somnoliencia:Somnolencia muscosidad:mucosidad agravacoin:agravación lechha:leche
respitatoria:respiratoria anrtes:antes ansidad:ansiedad profinda:profunda petroleun:petroleum leccha:leche
tendenciaa:tendencia intero:interno Trastronos:Trastornos membrsnas:membranas Patolgoia:Patología termnar:terminar
neceidad:necesidad curcubitacinos:cucurbitacinos Curcubitacinas:Cucurbitacinas cantitdades:cantidades Escaza:Escasa
dulen:duelen afinicdad:afinidad propencion:propensión iicial:inicial froma:forma lecha:leche precida:precedida
Transtornos:Trastornos transtornos:trastornos divion:división elecva:electiva totna:tonta vesicual:vesical
Tedencia:Tendencia afeccipnes:afecciones silimar:similar dolorsas:dolorosas HERMORRAGICAS:HEMORRÁGICAS inciio:inicio
sucucion:sucusión distencion:distensión intermentruales:intermenstruales exitacion:excitación sesancion:sensación
absesos:abscesos meticulocidad:meticulosidad acomulacion:acumulación etidas:fétidas puel:piel PIel:Piel
DElirio:Delirio atralgias:artralgias"""
def mapa(txt):
    return dict(p.split(":", 1) for p in txt.split())
TM, LM = mapa(T), mapa(L)

# Sustituciones por celda (texto exacto → texto); las de letras se declaran
CELDA = {
    "N2": [("agrava on calor", "agrava con calor"), ("filoso arapiento", "filósofo harapiento")],
    "G9": [("A2y", "A2 y")],
    "N9": [("por e sueño", "por el sueño")],
    "N17": [("enel", "en el")],
    "N19": [("lamusica", "la musica"), ("agrava po el", "agrava por el"), ("fracas,", "fracaso,")],
    "N21": [("contendencia", "con tendencia")],
    "N22": [("aliviojunto", "alivio junto")],
    "N27": [("fisurasy", "fisuras y")],
    "K28": [("picaen", "pica en")],
    "P29": [("Arnicadolor", "Arnica dolor")],
    "N31": [("lengua co saburra", "lengua con saburra"), ("y eseos", "y deseos")],
    "N32": [("frio,a l hablar", "frio, al hablar")],
    "P34": [("loq ue", "lo que"), ("solicito", "solicitó")],
    "N34": [("necesidad de ire", "necesidad de aire")],
    "H39": [("tuvier aun clavo", "tuviera un clavo"), ("Cefalea como si", "Cefalea como si")],
    "N42": [("no o deja", "no lo deja")],
    "N44": [("pot las", "por las"), ("co agua", "con agua")],
    "N38": [("imperioroy", "imperioso y"), ("a la plicaciones frial", "a las aplicaciones frías")],
    "L31": [("Sobretodo", "Sobre todo")],
    "K39": [("sobretodo", "sobre todo")],
    "B36": [("vulgares", "vulgaris")],
    "M22": [("esta indicado", "está indicado")],
    "I19": [("los mas importantes", "los más importantes")],
    "M27": [("graphites mas", "graphites más")],
    "P30": [("es mas periferico", "es más periférico")],
    "I40": [("no sabe que", "no sabe qué")],
    "L24": [("no sabe que tiene", "no sabe qué tiene")],
    "C23": [("aun no", "aún no")],
    "M31": [("este limpio", "esté limpio")],
    "M23": [("Cuando usar", "Cuándo usar")],
    "J39": [("vividos", "vívidos")],
    # saltos de línea que cortan una frase
    "C4": [("la\ncapa", "la capa")],
    "K4": [("de\narena", "de arena")],
    "G5": [("calcio\npueden", "calcio pueden")],
    "G22": [("IgA\nproteasa", "IgA proteasa")],
    "H22": [("blanco\namarillentas a\nmenudo", "blanco amarillentas a menudo"),
            ("Erupciones\nvesiculares\npruriginosas", "Erupciones vesiculares pruriginosas"),
            ("Esclerosis\ndel tejidos\nfibroconjuntivo que\nprovoca rigidez\ndolorosa",
             "Esclerosis del tejidos fibroconjuntivo que provoca rigidez dolorosa"),
            ("Comportamiento\nagitado", "Comportamiento agitado")],
    "I28": [("confusión y\ndesorientación", "confusión y desorientación"), ("calles\nque", "calles que"),
            ("está\ncerca", "está cerca"), ("sus\nasuntos", "sus asuntos"), ("se\nmolesta", "se molesta"),
            ("la\nmañana", "la mañana")],
    "H30": [("ricas en\nnervios:", "ricas en nervios:")],
    "I30": [("(mala\nmemoria)", "(mala memoria)"), ("post\ntraumática", "post traumática"),
            ("Hiperexcitabilidad\nnerviosa", "Hiperexcitabilidad nerviosa")],
    "J30": [("al\ncambio", "al cambio"), ("la\nhumedad, la niebla, el\ntacto y en espacios\ncerrados",
             "la humedad, la niebla, el tacto y en espacios cerrados"), ("la\ncabeza", "la cabeza")],
    "K30": [("compresión o\nde estallido", "compresión o de estallido"), ("Dolor\ndesproporcionado", "Dolor desproporcionado")],
    "L30": [("que\nafecten al sistema nervioso\ncentral", "que afecten al sistema nervioso central")],
    "G35": [("glotis\n(asfixia)", "glotis (asfixia)")],
    "H35": [("inflamación\naguda", "inflamación aguda"),
            ("Actúa sobre las\nmembranas que\nrecubren órganos\n(pleura, peritoneo,\nmeninges),\nprovocando derrames.",
             "Actúa sobre las membranas que recubren órganos (pleura, peritoneo, meninges), provocando derrames."),
            ("riñones\ny vejiga", "riñones y vejiga")],
    "I35": [("Deja\ncaer las cosas de\nlas manos.", "Deja caer las cosas de las manos."), ("Celos y\nsuspicacia", "Celos y suspicacia"),
            ("Grito encefálico\nen niños\n(especialmente\ncon fiebre alta o\nmeningitis).",
             "Grito encefálico en niños (especialmente con fiebre alta o meningitis)."),
            ("Es muy inquieto,\nse mueve mucho\npero no concreta\nnada.", "Es muy inquieto, se mueve mucho pero no concreta nada.")],
    "J35": [("Agravación por el\ncalor: No soporta el\ncalor de la\nhabitación, las\nmantas o el sol. Todo\nsíntoma mejora con\nel frío aplicado.",
             "Agravación por el calor: No soporta el calor de la habitación, las mantas o el sol. Todo síntoma mejora con el frío aplicado."),
            ("Dolor punzante como\nsi le estuvieran\nclavando agujas\ncalientes (como un\npicotazo de abeja).",
             "Dolor punzante como si le estuvieran clavando agujas calientes (como un picotazo de abeja).")],
    "K35": [("Hinchazón de\ncolor rosado o rojo\nbrillante. edema de los\npárpados inferiores,\nque cuelgan",
             "Hinchazón de color rosado o rojo brillante. edema de los párpados inferiores, que cuelgan"),
            ("Úvula\nhinchada", "Úvula hinchada"), ("ardor\ny picazón intensa que\nmejora", "ardor y picazón intensa que mejora"),
            ("difícil,\ncon goteo y dolor\nardiente (especialmente\nla última gota).",
             "difícil, con goteo y dolor ardiente (especialmente la última gota).")],
    "L35": [("insectos,\nurticarias y\nangioedemas.", "insectos, urticarias y angioedemas."),
            ("Sensacion de\ninflamación de la\nvejiga con mucho\nardor.", "Sensacion de inflamación de la vejiga con mucho ardor."),
            ("Inflamación\ndel ovario", "Inflamación del ovario"), ("rodilla con\ncalor.", "rodilla con calor.")],
    "H36": [("vejiga,\nmuslos", "vejiga, muslos")],
    "K36": [("vejiga,\nmuslos", "vejiga, muslos")],
    "I36": [("Irritabilidad\nsecundaria", "Irritabilidad secundaria"), ("para\nconcentrarse", "para concentrarse"),
            ("Sensación de\npesadez / lentitud\nmental", "Sensación de pesadez / lentitud mental")],
    "J36": [("momentáneamente el\ndolor", "momentáneamente el dolor")],
    "L36": [("(clave) ,\nHacia", "(clave) , Hacia"), ("Cambiante o\ncon irradiación", "Cambiante o con irradiación"),
            ("asociados\n:Alteraciones", "asociados :Alteraciones")],
    "I37": [("Tendencia a\ntristeza", "Tendencia a tristeza"), ("falta de\nmotivación", "falta de motivación"),
            ("Sensación de\n“opacidad”", "Sensación de “opacidad”")],
    "J37": [("Modalidades:no\npresenta", "Modalidades:no presenta")],
    "K37": [("irradiado a\ndientes", "irradiado a dientes"),
            ("Sensación de\n\n“banda de hierro”\n\nalrededor de la cabeza", "Sensación de “banda de hierro” alrededor de la cabeza")],
    "L37": [("(USO\nPRINCIPAL)", "(USO PRINCIPAL)"), ("supuración\nintensa", "supuración intensa")],
    "G43": [("altas\nsuprimen", "altas suprimen"), ("enrojecimiento\nocular", "enrojecimiento ocular")],
    "I43": [("segundos –\nsiglos", "segundos – siglos")],
    "K43": [("cabeza\nse abriera", "cabeza se abriera"), ("dolor\ndespués", "dolor después"), ("aumento\nesfuerzo", "aumento esfuerzo")],
    "J44": [("diferentes\npartes", "diferentes partes")],
    "K44": [("sobre\nel cráneo", "sobre el cráneo"), ("abundante y\nviscosa", "abundante y viscosa"),
            ("estando de\npie", "estando de pie"), ("lo\nhace", "lo hace"), ("Secreción\nuretral", "Secreción uretral")],
}

def caso(orig, nuevo):
    if orig.isupper() and len(orig) > 1:
        return nuevo.upper()
    if orig[0].isupper():
        return nuevo[0].upper() + nuevo[1:]
    return nuevo

correcciones = {}
def reemplazar_palabras(texto, mapa_, declarar):
    def f(m):
        w = m.group(0)
        clave = w if w in mapa_ else (w.lower() if w.lower() in mapa_ else None)
        if clave is None:
            return w
        nuevo = mapa_[clave] if clave == w and not w.islower() else caso(w, mapa_[clave])
        if declarar and st(nuevo).lower() != st(w).lower():
            correcciones[w] = nuevo
        return nuevo
    return re.sub(r"\w+", f, texto)

def mayuscula_inicial(texto):
    lineas = texto.split("\n")
    for i, l in enumerate(lineas):
        m = re.match(r"^((?:\s*(?:[-\*•]|\d+\s*[\.\-]|\(\+\)))*\s*)([^\W\d_])", l)
        if m and m.group(2).islower():
            lineas[i] = l[: m.start(2)] + m.group(2).upper() + l[m.end(2):]
    return "\n".join(lineas)

wb = openpyxl.load_workbook(DEST)
ws = wb.active
usados = set()
cambios = 0
for fila in ws.iter_rows():
    for c in fila:
        v = c.value
        if not isinstance(v, str):
            continue
        n = v
        for a, b in CELDA.get(c.coordinate, []):
            assert a in n, (c.coordinate, a)
            n = n.replace(a, b, 1)
            if re.sub(r"\s+", " ", st(a)).lower() != re.sub(r"\s+", " ", st(b)).lower():
                correcciones[a] = b
            usados.add(c.coordinate)
        if c.column_letter != "D":
            n = reemplazar_palabras(n, LM, True)
            n = reemplazar_palabras(n, TM, False)
            if c.coordinate != "C13":  # «a-bisabolol» (alfa), no es inicial de frase
                n = mayuscula_inicial(n)
        if n != v:
            c.value = n
            cambios += 1
faltan = set(CELDA) - usados
assert not faltan, faltan
assert ws["R33"].value == "#REF!"
ws["R33"].value = None
wb.save(DEST)
(BASE / ".trabajo/correcciones.json").write_text(json.dumps(correcciones, ensure_ascii=False, indent=1), "utf-8")
print("celdas cambiadas", cambios, "correcciones", len(correcciones))
