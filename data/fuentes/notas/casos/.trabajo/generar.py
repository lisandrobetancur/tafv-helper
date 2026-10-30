"""Genera notas/*.md y plan.json del componente Casos a partir del Word (y verifica el render)."""
import json
import re
import sys
import unicodedata
from pathlib import Path

import docx
from markdown_it import MarkdownIt

BASE = Path(__file__).resolve().parent.parent
DOC = BASE / "documentos" / "Notas casos.docx"
_md = MarkdownIt("commonmark", {"html": False, "breaks": True, "typographer": True}).enable(["table", "strikethrough"])

INICIOS = [1, 79, 158, 238, 325, 494, 564, 590, 608, 619, 645, 672, 690, 711, 736, 770, 839, 888, 931, 978,
           1002, 1026, 1052, 1078, 1119, 1148]

# palabra original -> corregida (solo cambios de letras seguros)
CORRECCIONES = {
    "FISCO": "FÍSICO",
    "magistrarl": "magistral",
    "exposiocion": "exposición",
    "minitos": "minutos",
    "procedinmiento": "procedimiento",
    "Autohemoteraoia": "Autohemoterapia",
    "sculentum": "esculentum",
    "Paullina": "Paullinia",
    "Nerural": "Neural",
    "Colistis": "Colitis",
    "SULHUR": "SULPHUR",
    "PHOSPUROS": "PHOSPHORUS",
    "CAUTISCUM": "CAUSTICUM",
    "sulfhur": "sulphur",
    "phosphoros": "phosphorus",
    "pulsatalia": "pulsatilla",
    "folliculim": "folliculinum",
    "Calcalera": "Calcarea",
    "metalicum": "metallicum",
    "Hodking": "Hodgkin",
    "Pegfilgrastrim": "Pegfilgrastim",
    "Alupurinol": "Alopurinol",
    "Hyosciamus": "Hyoscyamus",
    "mollisima": "mollissima",
    "cebo": "sebo",
}
_re_corr = re.compile(r"(?<![\wÀ-ÿ])(" + "|".join(map(re.escape, CORRECCIONES)) + r")(?![\wÀ-ÿ])")
usadas: dict[str, int] = {}


def corregir(t: str) -> str:
    def f(m):
        usadas[m.group(1)] = usadas.get(m.group(1), 0) + 1
        return CORRECCIONES[m.group(1)]
    return _re_corr.sub(f, t)


def esc(t: str) -> str:
    return t.replace("\\", "\\\\").replace("*", "\\*").replace("_", "\\_").replace("`", "\\`")


def sin_tildes(t):
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")


SECCION = re.compile(
    r"^(caso clinico|motivo de consulta|enfermedad actual|revision por sistemas?\.?( \(-\))?|"
    r"antecedentes( personales| personal)?[.:]?|(iv\. )?examen (fisico|fisco):?|diagnosticos?:?|tratamiento:?|"
    r"jerarquizacion:?|idx:|analisis( del caso)?:|evolucion:|sintomas considerados para repertorizacion:)$")


def es_seccion(t: str) -> bool:
    return bool(SECCION.match(" ".join(sin_tildes(t).lower().split())))


LISTA = re.compile(r"^(\d+)([.)])(\s|$)")
ROMANO = "abcdefghijklmnopqrstuvwxyz"


def bold_runs(p):
    runs = [r for r in p.runs if r.text]
    con_texto = [r for r in runs if r.text.strip()]
    if not con_texto:
        return "ninguno", runs
    bs = [bool(r.bold) for r in con_texto]
    return ("todo" if all(bs) else "parte" if any(bs) else "ninguno"), runs


def con_negritas(runs) -> str:
    """Texto con **…** en los tramos en negrita (espacios fuera de las marcas)."""
    grupos = []
    for r in runs:
        b = bool(r.bold) and bool(r.text.strip()) or (bool(r.bold) and grupos and grupos[-1][0])
        if grupos and grupos[-1][0] == b:
            grupos[-1][1] += r.text
        else:
            grupos.append([b, r.text])
    salida = ""
    for b, t in grupos:
        t = esc(corregir(t))
        if b:
            # cada renglón por separado; espacios fuera de las marcas
            partes = []
            for ln in t.split("\n"):
                m = re.match(r"^(\s*)(.*?)(\s*)$", ln, re.S)
                partes.append(f"{m.group(1)}**{m.group(2)}**{m.group(3)}" if m.group(2) else ln)
            t = "\n".join(partes)
        salida += t
    return salida


def generar(sin_negrita: set[int], escapar: set[tuple[int, int]]):
    d = docx.Document(str(DOC))
    ps = d.paragraphs
    notas = []
    fines = [i - 1 for i in INICIOS[1:]] + [len(ps)]
    for n_caso, (ini, fin) in enumerate(zip(INICIOS, fines), start=1):
        assert re.match(r"^\s*CASO\s+%d\s*$" % n_caso, ps[ini - 1].text), (ini, ps[ini - 1].text)
        lineas = []   # (tipo, texto, parrafo)   tipo: blank | head | item | sub | text
        cont27 = [0, 0]
        for i in range(ini + 1, fin + 1):
            p = ps[i - 1]
            texto = p.text
            if not texto.strip():
                lineas.append(("blank", "", i))
                continue
            negrita, runs = bold_runs(p)
            pPr = p._p.pPr
            numpr = pPr.numPr if pPr is not None else None
            if numpr is not None:
                nid = numpr.numId.val
                lvl = numpr.ilvl.val if numpr.ilvl is not None else 0
                t = esc(corregir(texto.strip()))
                if nid == 25:
                    lineas.append(("item", f"- {t}", i))
                elif nid == 27:
                    if lvl == 0:
                        cont27[0] += 1
                        cont27[1] = 0
                        lineas.append(("item", f"{cont27[0]}. {t}", i))
                    else:
                        cont27[1] += 1
                        lineas.append(("sub", f"   {ROMANO[cont27[1] - 1]}. {t}", i))
                else:
                    raise SystemExit(f"numeración desconocida {nid} en {i}")
                continue
            if (negrita == "todo" or es_seccion(texto)) and "\n" not in texto.strip():
                lineas.append(("head", "### " + esc(corregir(texto.strip())), i))
                continue
            if negrita == "parte" and i not in sin_negrita:
                bruto = con_negritas(runs)
            else:
                bruto = esc(corregir(texto))
            for k, ln in enumerate(bruto.split("\n")):
                s = ln.strip()
                if s.startswith("• "):
                    lineas.append(("item", "- " + s[2:].strip(), i))
                elif LISTA.match(s):
                    if (i, k) in escapar:
                        m = LISTA.match(s)
                        lineas.append(("text", f"{m.group(1)}\\{m.group(2)}{s[len(m.group(1)) + 1:]}", i))
                    else:
                        lineas.append(("item", s, i, k))
                else:
                    lineas.append(("text", s, i))
        # armar el Markdown
        out = []
        previo = None
        for ln in lineas:
            tipo, t = ln[0], ln[1]
            if tipo == "blank":
                if out and out[-1] != "":
                    out.append("")
                previo = "blank"
                continue
            if tipo == "head":
                if out and out[-1] != "":
                    out.append("")
                out.append(t)
                out.append("")
                previo = "head"
                continue
            if tipo == "text" and previo in ("item", "sub") and out and out[-1] != "":
                out.append("")
            if tipo == "item" and previo == "text" and out and out[-1] != "" and not t.startswith(("- ",)) and not LISTA.match(t):
                out.append("")
            out.append(t)
            previo = tipo
        while out and out[-1] == "":
            out.pop()
        md = "\n".join(out) + "\n"
        notas.append({"n": n_caso, "ini": ini, "fin": fin, "md": md, "lineas": lineas})
    return notas


def visibles(md: str) -> list[str]:
    """Palabras que se ven en el HTML renderizado, con los números que muestran las listas."""
    tokens = _md.parse(md)
    pila = []
    salida = []
    for tk in tokens:
        if tk.type == "ordered_list_open":
            pila.append(["ol", int(tk.attrGet("start") or 1) - 1])
        elif tk.type == "bullet_list_open":
            pila.append(["ul", 0])
        elif tk.type in ("ordered_list_close", "bullet_list_close"):
            pila.pop()
        elif tk.type == "list_item_open":
            if pila[-1][0] == "ol":
                pila[-1][1] += 1
                salida.append(f"{pila[-1][1]}{tk.markup}")
            else:
                salida.append("•")
        elif tk.type == "inline":
            texto = ""
            for ch in tk.children:
                if ch.type in ("text", "code_inline"):
                    texto += ch.content
                elif ch.type in ("softbreak", "hardbreak"):
                    texto += " "
            salida.extend(texto.split())
    return salida


def esperado(ps, ini, fin) -> list[str]:
    salida = []
    cont27 = [0, 0]
    for i in range(ini + 1, fin + 1):
        p = ps[i - 1]
        t = corregir_silencioso(p.text)
        pPr = p._p.pPr
        if pPr is not None and pPr.numPr is not None and t.strip():
            nid, lvl = pPr.numPr.numId.val, (pPr.numPr.ilvl.val if pPr.numPr.ilvl is not None else 0)
            if nid == 25:
                salida.append("•")
            elif lvl == 0:
                cont27[0] += 1
                cont27[1] = 0
                salida.append(f"{cont27[0]}.")
            else:
                cont27[1] += 1
                salida.append(f"{ROMANO[cont27[1] - 1]}.")
        salida.extend(t.split())
    return salida


def corregir_silencioso(t):
    return _re_corr.sub(lambda m: CORRECCIONES[m.group(1)], t)


def main():
    d = docx.Document(str(DOC))
    ps = d.paragraphs
    sin_negrita: set[int] = set()
    escapar: set[tuple[int, int]] = set()
    for vuelta in range(10):
        usadas.clear()
        notas = generar(sin_negrita, escapar)
        cambios = False
        for nt in notas:
            tokens = _md.parse(nt["md"])
            # números de lista mostrados vs. los de la fuente
            mdlineas = nt["md"].split("\n")
            pila = []
            for tk in tokens:
                if tk.type == "ordered_list_open":
                    pila.append([int(tk.attrGet("start") or 1) - 1, [], False])
                elif tk.type == "ordered_list_close":
                    _, items, malo = pila.pop()
                    if malo:   # si un número no coincide, se escapa toda la lista
                        for linea in items:
                            for ln in nt["lineas"]:
                                if len(ln) == 4 and ln[1] == linea:
                                    escapar.add((ln[2], ln[3]))
                                    cambios = True
                elif tk.type == "bullet_list_open":
                    pila.append(None)
                elif tk.type == "bullet_list_close":
                    pila.pop()
                elif tk.type == "list_item_open" and pila[-1] is not None:
                    pila[-1][0] += 1
                    linea = mdlineas[tk.map[0]].strip()
                    pila[-1][1].append(linea)
                    m = LISTA.match(linea)
                    if m and int(m.group(1)) != pila[-1][0]:
                        pila[-1][2] = True
            # negritas que no se renderizan
            vis = visibles(nt["md"])
            if any("**" in w for w in vis):
                for ln in nt["lineas"]:
                    if "**" in ln[1]:
                        r = visibles(ln[1])
                        if any("**" in w for w in r):
                            sin_negrita.add(ln[2])
                            cambios = True
        if not cambios:
            break
    # verificación final: lo visible == la fuente, palabra por palabra
    errores = 0
    for nt in notas:
        vis = [w for w in visibles(nt["md"])]
        esp = esperado(ps, nt["ini"], nt["fin"])
        norm = lambda ws: [w.replace("•", "•") for w in ws if w != "•" or True]
        # las viñetas «•» de la fuente se ven como viñeta de lista
        if norm(vis) != norm(esp):
            errores += 1
            import difflib
            print(f"CASO {nt['n']}: diferencia visible/fuente")
            for l in difflib.unified_diff(esp, vis, lineterm="", n=2):
                print("   ", l)
    print("escapados:", sorted(escapar))
    print("sin negrita:", sorted(sin_negrita))
    print("correcciones usadas:", usadas)
    faltan = set(CORRECCIONES) - set(usadas)
    if faltan:
        print("CORRECCIONES SIN USO:", faltan)
    if "--escribir" in sys.argv:
        (BASE / "notas").mkdir(exist_ok=True)
        for f in (BASE / "notas").glob("*.md"):
            f.unlink()
        plan_notas = []
        for nt in notas:
            titulo = f"CASO {nt['n']} - Casos"
            archivo = f"notas/{titulo}.md"
            (BASE / archivo).write_text(nt["md"], encoding="utf-8")
            plan_notas.append({"titulo": titulo, "archivo": archivo, "parrafos": f"{nt['ini']}-{nt['fin']}"})
        plan = {
            "materia": "Notas",
            "componente": "Casos",
            "crear_componente": {"tipos": ["notas"], "descripcion": ""},
            "fuentes": ["documentos/Notas casos.docx"],
            "correcciones": {k: v for k, v in CORRECCIONES.items()},
            "omitidos": [],
            "notas": plan_notas,
            "tablas": [],
            "imagenes": [],
            "pdfs": [],
        }
        (BASE / "plan.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print("escrito")
    sys.exit(1 if errores else 0)


if __name__ == "__main__":
    main()
