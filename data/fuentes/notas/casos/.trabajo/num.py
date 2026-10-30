import docx
from docx.oxml.ns import qn
d = docx.Document("data/fuentes/notas/casos/documentos/Notas casos.docx")
nb = d.part.numbering_part.element
for num in nb.findall(qn('w:num')):
    nid = num.get(qn('w:numId'))
    if nid not in ('25','27'): continue
    aid = num.find(qn('w:abstractNumId')).get(qn('w:val'))
    print("num", nid, "abstract", aid, [o.tag for o in num])
    for a in nb.findall(qn('w:abstractNum')):
        if a.get(qn('w:abstractNumId')) == aid:
            for lvl in a.findall(qn('w:lvl'))[:3]:
                g = lambda t: (lvl.find(qn(t)).get(qn('w:val')) if lvl.find(qn(t)) is not None else None)
                print("  lvl", lvl.get(qn('w:ilvl')), g('w:start'), g('w:numFmt'), repr(g('w:lvlText')))
# partial bold detail
for i, p in enumerate(d.paragraphs, 1):
    runs = [r for r in p.runs if r.text.strip()]
    bs = [bool(r.bold) for r in runs]
    if runs and any(bs) and not all(bs):
        print(i, [(r.text, bool(r.bold)) for r in p.runs])
# characters check
import re
for i, p in enumerate(d.paragraphs, 1):
    t = p.text
    for ln in t.split("\n"):
        s = ln.strip()
        if re.match(r'^([-+*>#|=`~]|\d+[.)]\s)', s) or re.search(r'[<_`\\|]|\(c\)|\(r\)|\(tm\)|\(p\)|\+-|\.\.\.|--|,,|!!!!|\?\?\?\?|\t', ln, re.I):
            print("CHK", i, repr(ln))
